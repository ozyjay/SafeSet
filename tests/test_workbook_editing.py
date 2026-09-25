"""Synthetic editable-workbook round trips and adversarial restoration checks."""

import copy
import json
import socket
import zipfile
from types import SimpleNamespace
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table as ExcelTable

from safeset.desktop_bridge import Bridge
from safeset.desktop_flow import (
    approve_relational_protection,
    approve_relational_reconstruction,
    prepare_relational_protection,
    prepare_relational_reconstruction,
)
from safeset.errors import SafetyError
from safeset.ingestion import Table, excel_workbook_bytes, read_excel_sheets
from safeset.policy_authoring import RuleDraft
from safeset.relational import (
    decrypt_relational_bundle,
    encrypt_relational_bundle,
    read_relational_bundle,
    review_relational_reconstruction,
    validate_relational_bundle,
)

from .conftest import PASSPHRASE
from .test_desktop_bridge import call


@pytest.fixture
def editing(tmp_path, destinations):
    source = tmp_path / "synthetic-original.xlsx"
    workbook = Workbook()
    allocations = workbook.active
    allocations.title = "Allocations"
    for row in [
        ("student_key", "team", "state", "score"),
        ("SYNTH-001", "Synthetic Alpha", "Current", None),
        ("SYNTH-002", "Synthetic Alpha", "Pending", 2),
    ]:
        allocations.append(row)
    allocations.move_range("A1:D3", rows=2, cols=1)
    allocations["A1"] = "Synthetic private title"
    allocations.merge_cells("A1:D1")
    allocations["C4"].fill = PatternFill("solid", fgColor="FFFF00")
    allocations.add_table(ExcelTable(displayName="SyntheticAllocations", ref="B3:E5"))
    participants = workbook.create_sheet("Participants")
    participants.append(("student_key", "team"))
    for key, team in [
        ("SYNTH-001", "Synthetic Alpha"),
        ("SYNTH-002", "Synthetic Alpha"),
        ("SYNTH-003", "Synthetic Beta"),
    ]:
        participants.append((key, team))
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Synthetic local summary"
    summary.merge_cells("A1:D1")
    summary["A2"] = "=SUM(Allocations!E4:E5)"
    summary["B2"] = 2
    chart = BarChart()
    chart.add_data(Reference(summary, min_col=2, min_row=2, max_row=2))
    summary.add_chart(chart, "F1")
    workbook.save(source)
    drafts = {
        "Allocations": {
            "student_key": RuleDraft("pseudonymise", "direct_identifier"),
            "team": RuleDraft("code", "analytical_attribute", ("Synthetic Alpha",)),
            "state": RuleDraft("keep", "analytical_attribute", ("Current", "Pending")),
            "score": RuleDraft("keep_numeric", "analytical_attribute", bounds=(0, 10)),
        },
        "Participants": {
            "student_key": RuleDraft("pseudonymise", "direct_identifier"),
            "team": RuleDraft(
                "code", "analytical_attribute", ("Synthetic Alpha", "Synthetic Beta")
            ),
        },
    }
    fields = {"Allocations": ["team", "state", "score"], "Participants": []}
    output, bundle = destinations
    review = prepare_relational_protection(
        source,
        tuple(drafts),
        drafts,
        "2",
        output,
        bundle,
        "controlled_pseudonymisation",
        ("team",),
        fields,
    )
    approve_relational_protection(review, PASSPHRASE, approved=True)
    returned = tmp_path / "synthetic-returned.xlsx"
    returned.write_bytes(output.read_bytes())
    return SimpleNamespace(
        source=source,
        protected=output,
        bundle=bundle,
        returned=returned,
        restored=tmp_path / "synthetic-restored.xlsx",
        drafts=drafts,
        fields=fields,
        protection=review,
    )


def returned_tables(files):
    return read_excel_sheets(files.returned, tuple(files.drafts))


def review(files):
    return prepare_relational_reconstruction(
        files.returned,
        files.source,
        files.bundle,
        files.restored,
        PASSPHRASE,
    )


def change_allocation(files):
    tables = returned_tables(files)
    allocations = tables["Allocations"]
    rows = [dict(row) for row in allocations.rows]
    rows[0].update(team=tables["Participants"].rows[2]["team"], state="Pending", score="4")
    tables["Allocations"] = Table(allocations.columns, tuple(reversed(rows)))
    files.returned.write_bytes(excel_workbook_bytes(tables))


def test_edit_round_trip_preserves_reference_parts_styles_formulas_and_source_order(
    editing,
    monkeypatch,
):
    def denied(*_args, **_kwargs):
        raise AssertionError("Network access attempted")

    monkeypatch.setattr(socket, "socket", denied)
    monkeypatch.setattr(socket, "create_connection", denied)
    change_allocation(editing)
    prepared = review(editing)
    assert prepared.changes == {
        "Allocations": {"team": 1, "state": 1, "score": 1},
        "Participants": {},
    }
    assert not editing.restored.exists()
    with pytest.raises(SafetyError, match="changed field"):
        approve_relational_reconstruction(prepared, prepared.new_columns, authorised=True)
    assert not editing.restored.exists()
    approve_relational_reconstruction(
        prepared,
        prepared.new_columns,
        authorised=True,
        approved_changes={"Allocations": ("team", "state", "score"), "Participants": ()},
    )
    restored = load_workbook(editing.restored)
    assert restored["Allocations"]["B4"].value == "SYNTH-001"
    assert restored["Allocations"]["C4"].value == "Synthetic Beta"
    assert restored["Allocations"]["D4"].value == "Pending"
    assert restored["Allocations"]["E4"].value == 4
    assert restored["Allocations"]["E4"].data_type == "n"
    assert restored["Allocations"]["C4"].fill.fgColor.rgb == "00FFFF00"
    assert restored["Allocations"]["C5"].value == "Synthetic Alpha"
    assert restored["Summary"]["A2"].value == "=SUM(Allocations!E4:E5)"
    assert restored.calculation.fullCalcOnLoad
    assert restored.calculation.forceFullCalc
    with zipfile.ZipFile(editing.source) as original, zipfile.ZipFile(editing.restored) as result:
        assert original.namelist() == result.namelist()
        changed = {name for name in original.namelist() if original.read(name) != result.read(name)}
        assert changed == {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
    original = load_workbook(editing.source)
    assert original["Allocations"]["C4"].value == "Synthetic Alpha"
    assert original["Allocations"]["E4"].value is None
    protected = str(returned_tables(editing))
    assert "SYNTH-001" not in protected and "Synthetic Beta" not in protected


@pytest.mark.parametrize(
    "mutation",
    [
        "unknown_code",
        "blank_code",
        "numeric_bound",
        "category",
        "readonly",
        "entity",
        "missing",
        "duplicate",
        "unknown_id",
        "extra_field",
        "extra_sheet",
        "formula",
        "hidden",
    ],
)
def test_unapproved_changes_fail_closed(editing, mutation):
    workbook = load_workbook(editing.returned)
    sheet = workbook["Allocations"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    if mutation == "unknown_code":
        sheet.cell(2, columns["team"], "Synthetic Beta")
    elif mutation == "blank_code":
        sheet.cell(2, columns["team"]).value = None
    elif mutation == "numeric_bound":
        sheet.cell(2, columns["score"], "11")
    elif mutation == "category":
        sheet.cell(2, columns["state"], "Unknown")
    elif mutation == "readonly":
        workbook["Participants"]["C2"] = workbook["Participants"]["C4"].value
    elif mutation == "entity":
        sheet["B2"] = sheet["B3"].value
    elif mutation == "missing":
        sheet.delete_rows(2)
    elif mutation == "duplicate":
        sheet["A2"] = sheet["A3"].value
    elif mutation == "unknown_id":
        sheet["A2"] = "synthetic-unknown-id"
    elif mutation == "extra_field":
        sheet.cell(1, 7, "Finding")
        sheet.cell(2, 7, "Synthetic")
    elif mutation == "extra_sheet":
        workbook.create_sheet("Findings").append(("Note",))
    elif mutation == "formula":
        sheet.cell(2, columns["score"], "=2+2")
    else:
        workbook["Participants"].sheet_state = "hidden"
    workbook.save(editing.returned)
    with pytest.raises(SafetyError):
        review(editing)
    assert not editing.restored.exists()


def test_unchanged_editing_workbook_is_an_exact_copy(editing):
    prepared = review(editing)
    assert prepared.changes == {"Allocations": {}, "Participants": {}}
    approve_relational_reconstruction(
        prepared,
        prepared.new_columns,
        authorised=True,
        approved_changes={"Allocations": (), "Participants": ()},
    )
    assert editing.restored.read_bytes() == editing.source.read_bytes()


@pytest.mark.parametrize("where", ["before_review", "after_review", "returned_after_review"])
def test_full_source_and_returned_changes_invalidate_review(editing, where):
    prepared = review(editing) if where != "before_review" else None
    path = editing.returned if where == "returned_after_review" else editing.source
    workbook = load_workbook(path)
    if path == editing.source:
        workbook["Summary"]["B2"] = 3  # Unselected local content is also bound.
    else:
        workbook["Allocations"]["E2"] = 4
    workbook.save(path)
    with pytest.raises(SafetyError):
        if prepared:
            approve_relational_reconstruction(
                prepared,
                prepared.new_columns,
                authorised=True,
                approved_changes={"Allocations": (), "Participants": ()},
            )
        else:
            review(editing)
    assert not editing.restored.exists()


@pytest.mark.parametrize(
    "fields",
    [
        {},
        {"Allocations": ["student_key"], "Participants": []},
        {"Allocations": ["team", "team"], "Participants": []},
        {"Allocations": ["missing"], "Participants": []},
        {"Allocations": "team", "Participants": []},
        {"Allocations": [True], "Participants": []},
    ],
)
def test_edit_permissions_are_strict_and_versioned(editing, fields):
    bundle = read_relational_bundle(editing.bundle, PASSPHRASE)
    bundle["editable_fields"] = fields
    with pytest.raises(SafetyError):
        validate_relational_bundle(bundle)


def test_bundle_version_migration_does_not_grant_edits_to_old_bundles(editing):
    bundle = read_relational_bundle(editing.bundle, PASSPHRASE)
    assert bundle["version"] == 4
    assert bundle["editable_fields"] == editing.fields
    assert (
        decrypt_relational_bundle(encrypt_relational_bundle(bundle, PASSPHRASE), PASSPHRASE)
        == bundle
    )
    with pytest.raises(SafetyError):
        decrypt_relational_bundle(b"SAFESET3\n" + editing.bundle.read_bytes()[9:], PASSPHRASE)
    for key, value in [("unknown", []), ("source_file_digest", True)]:
        malformed = {**bundle, key: value}
        with pytest.raises(SafetyError):
            validate_relational_bundle(malformed)
    old = copy.deepcopy(bundle)
    old["version"] = 3
    del old["editable_fields"], old["source_file_digest"]
    validate_relational_bundle(old)
    change_allocation(editing)
    with pytest.raises(SafetyError, match="protected source field"):
        review_relational_reconstruction(
            editing.protection.source_tables,
            returned_tables(editing),
            old,
        )


def test_bridge_reports_counts_and_requires_exact_change_approval(editing):
    change_allocation(editing)
    bridge = Bridge()
    payload = {
        "returned": str(editing.returned),
        "source": str(editing.source),
        "bundle": str(editing.bundle),
        "output": str(editing.restored),
        "passphrase": PASSPHRASE,
    }
    response = call(bridge, "prepare_relational_reconstruction", payload)
    assert response["ok"]
    encoded = json.dumps(response)
    assert "SYNTH-001" not in encoded and "Synthetic Beta" not in encoded
    assert PASSPHRASE not in encoded
    prepared = response["result"]
    assert prepared["changes"]["Allocations"]["team"] == 1
    rejected = call(
        bridge,
        "approve_relational_reconstruction",
        {
            "review_id": prepared["review_id"],
            "approved_results": prepared["new_columns"],
            "approved_changes": {"Allocations": ["team"], "Participants": []},
        },
    )
    assert rejected["error"] == "edit_approval"
    assert not editing.restored.exists()
    prepared = call(bridge, "prepare_relational_reconstruction", payload)["result"]
    response = call(
        bridge,
        "approve_relational_reconstruction",
        {
            "review_id": prepared["review_id"],
            "approved_results": prepared["new_columns"],
            "approved_changes": {"Allocations": ["team", "state", "score"], "Participants": []},
        },
    )
    assert response["ok"] and editing.restored.exists()


def test_bridge_protection_binds_permissions_and_provides_a_value_free_prompt(editing):
    bridge = Bridge()
    payload = {
        "source": str(editing.source),
        "sheets": list(editing.drafts),
        "output": str(editing.protected.with_name("synthetic-second.xlsx")),
        "bundle": str(editing.bundle.with_name("synthetic-second.enc")),
        "threshold": "2",
        "validation_profile": "controlled_pseudonymisation",
        "shared_code_fields": ["team"],
        "editable_fields": editing.fields,
        "drafts": {
            sheet: {
                name: {
                    "action": draft.action,
                    "classification": draft.classification,
                    "allowed_values": list(draft.allowed_values),
                    "bins": [],
                    "bounds": list(draft.bounds) if draft.bounds else None,
                }
                for name, draft in drafts.items()
            }
            for sheet, drafts in editing.drafts.items()
        },
    }
    response = call(bridge, "prepare_relational_protection", payload)
    assert response["ok"]
    result = response["result"]
    assert result["editable_fields"] == editing.fields
    prompt = result["analysis_prompt"]
    assert "SYNTH-001" not in prompt and "Synthetic Alpha" not in prompt
    assert "from 0 to 10" in prompt and "record_id" in prompt
    assert '"Participants": {}' in prompt
    assert not editing.bundle.with_name("synthetic-second.enc").exists()
    response = call(
        bridge,
        "approve_relational_protection",
        {
            "review_id": result["review_id"],
            "passphrase": PASSPHRASE,
        },
    )
    assert response["ok"]
    assert editing.bundle.with_name("synthetic-second.enc").read_bytes().startswith(b"SAFESET4\n")


def test_numeric_normalisation_does_not_create_a_spurious_change(editing):
    workbook = load_workbook(editing.returned)
    workbook["Allocations"]["E3"] = "02.000"
    workbook.save(editing.returned)
    prepared = review(editing)
    assert prepared.changes == {"Allocations": {}, "Participants": {}}
    assert prepared.workbook_bytes == editing.source.read_bytes()


def test_numeric_precision_is_checked_before_user_approval(editing):
    workbook = load_workbook(editing.returned)
    workbook["Allocations"]["E2"] = "0.1234567890123456"
    workbook.save(editing.returned)
    with pytest.raises(SafetyError, match="precision"):
        review(editing)
    assert not editing.restored.exists()


def test_formula_fields_cannot_be_authorised_for_editing(editing):
    workbook = load_workbook(editing.source)
    workbook["Allocations"]["E4"] = "=2+2"
    workbook.save(editing.source)
    with zipfile.ZipFile(editing.source) as archive:
        entries = [(entry, archive.read(entry)) for entry in archive.infolist()]
    with zipfile.ZipFile(editing.source, "w") as archive:
        for entry, data in entries:
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                cell = root.find(f".//{ns}c[@r='E4']")
                cached = cell.find(f"{ns}v")
                if cached is None:
                    cached = ElementTree.SubElement(cell, f"{ns}v")
                cached.text = "4"
                data = ElementTree.tostring(root)
            archive.writestr(entry, data)
    with pytest.raises(SafetyError, match="Formula and date"):
        prepare_relational_protection(
            editing.source,
            tuple(editing.drafts),
            editing.drafts,
            "2",
            editing.protected.with_name("synthetic-formula.xlsx"),
            editing.bundle.with_name("synthetic-formula.enc"),
            "controlled_pseudonymisation",
            ("team",),
            editing.fields,
        )


def test_source_package_changes_block_protection_publication(editing):
    prepared = prepare_relational_protection(
        editing.source,
        tuple(editing.drafts),
        editing.drafts,
        "2",
        editing.protected.with_name("synthetic-stale.xlsx"),
        editing.bundle.with_name("synthetic-stale.enc"),
        "controlled_pseudonymisation",
        ("team",),
        editing.fields,
    )
    workbook = load_workbook(editing.source)
    workbook["Summary"]["B2"] = 7
    workbook.save(editing.source)
    with pytest.raises(SafetyError, match="changed after"):
        approve_relational_protection(prepared, PASSPHRASE, approved=True)
    assert not prepared.output.exists() and not prepared.bundle_path.exists()
