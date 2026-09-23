import stat

import pytest
from openpyxl import load_workbook

from safeset.errors import SafetyError
from safeset.ingestion import Table, excel_bytes, read_excel
from safeset.policy import load_policy
from safeset.policy_authoring import (
    RuleDraft,
    load_drafts,
    local_categories,
    local_category_review,
    parse_number,
    parse_pairs,
    policy_payload,
    save_policy,
)
from safeset.transform import sanitise
from safeset.validation import validate

from .conftest import ROOT


def example_drafts():
    return {
        "student_name": RuleDraft("drop", "direct_identifier"),
        "student_number": RuleDraft("pseudonymise", "direct_identifier"),
        "email": RuleDraft("drop", "direct_identifier"),
        "campus": RuleDraft("code", "quasi_identifier", ("Moon", "Mars")),
        "subject": RuleDraft("code", "quasi_identifier", ("Imaginary101",)),
        "gpa": RuleDraft("keep_numeric", "quasi_identifier", bounds=(0, 7)),
        "notes": RuleDraft("drop", "free_text"),
    }


def test_policy_builder_saves_private_strict_policy(tmp_path):
    destination = tmp_path / "new-policy.yaml"
    saved = save_policy(
        ROOT / "examples/synthetic_students.xlsx", destination, example_drafts(), "2"
    )
    policy = load_policy(saved)
    assert policy.version == 3
    assert policy.columns["campus"].action == "code"
    assert policy.columns["gpa"].action == "keep_numeric"
    assert stat.S_IMODE(saved.stat().st_mode) == 0o600
    assert "SYNTH-001" not in saved.read_text()
    assert validate(
        sanitise(read_excel(ROOT / "examples/synthetic_students.xlsx"), policy).table, policy
    ).passed
    with pytest.raises(SafetyError, match="overwriting"):
        save_policy(ROOT / "examples/synthetic_students.xlsx", destination, example_drafts(), "2")


def test_removed_field_needs_no_classification_choice():
    drafts = example_drafts()
    drafts["student_name"].classification = ""
    payload = policy_payload(drafts, "2")
    assert payload["columns"]["student_name"] == {
        "action": "drop",
        "classification": "unknown",
    }
    assert payload["columns"]["email"]["classification"] == "direct_identifier"


def test_policy_builder_uses_literal_source_headings(tmp_path):
    source = tmp_path / "synthetic.xlsx"
    source.write_bytes(
        excel_bytes(
            Table(
                ("Student Number", "Campus Location"),
                (
                    {"Student Number": "SYNTH-001", "Campus Location": "Moon"},
                    {"Student Number": "SYNTH-002", "Campus Location": "Moon"},
                ),
            )
        )
    )
    drafts = {
        "Student Number": RuleDraft("pseudonymise", "direct_identifier"),
        "Campus Location": RuleDraft("keep", "quasi_identifier", ("Moon",)),
    }
    saved = save_policy(source, tmp_path / "policy.yaml", drafts, "2")
    policy = load_policy(saved)
    assert tuple(policy.columns) == ("Student Number", "Campus Location")
    assert validate(sanitise(read_excel(source), policy).table, policy).passed


def test_existing_policy_can_be_loaded_and_saved_as_new_file(tmp_path):
    source = ROOT / "examples/synthetic_students.xlsx"
    drafts, threshold = load_drafts(source, ROOT / "examples/example-policy.yaml")
    assert drafts["campus"].action == "code"
    assert drafts["gpa"].bounds == (0, 7)
    saved = save_policy(source, tmp_path / "revised.yaml", drafts, str(threshold))
    migrated = load_policy(saved)
    assert migrated.version == 3
    assert migrated.columns["gpa"].action == "keep_numeric"


def test_policy_authoring_uses_selected_worksheet(tmp_path):
    source = tmp_path / "multiple.xlsx"
    workbook = load_workbook(ROOT / "examples/synthetic_students.xlsx")
    workbook.active.title = "Allocations"
    workbook.create_sheet("Instructions").append(("Invented help",))
    workbook.save(source)
    drafts, threshold = load_drafts(source, ROOT / "examples/example-policy.yaml", "Allocations")
    assert local_categories(source, "campus", "Allocations") == ("Mars", "Moon")
    saved = save_policy(source, tmp_path / "selected.yaml", drafts, str(threshold), "Allocations")
    assert load_policy(saved).columns["campus"].action == "code"


@pytest.mark.parametrize("threshold", ["", "1", "two", "2.0"])
def test_policy_builder_rejects_bad_threshold_before_writing(tmp_path, threshold):
    destination = tmp_path / "invalid.yaml"
    with pytest.raises(SafetyError):
        save_policy(
            ROOT / "examples/synthetic_students.xlsx", destination, example_drafts(), threshold
        )
    assert not destination.exists()


def test_policy_builder_rejects_unclassified_field_and_schema_drift(tmp_path):
    destination = tmp_path / "invalid.yaml"
    drafts = example_drafts()
    drafts["campus"].classification = ""
    with pytest.raises(SafetyError):
        save_policy(ROOT / "examples/synthetic_students.xlsx", destination, drafts, "2")
    drafts = example_drafts()
    drafts.pop("notes")
    with pytest.raises(SafetyError, match="headings changed"):
        save_policy(ROOT / "examples/synthetic_students.xlsx", destination, drafts, "2")
    assert not destination.exists()


def test_policy_builder_rejects_repository_destination():
    destination = ROOT / "unsafe-operational-policy.yaml"
    with pytest.raises(SafetyError, match="outside repositories"):
        save_policy(ROOT / "examples/synthetic_students.xlsx", destination, example_drafts(), "2")
    assert not destination.exists()


def test_policy_builder_numeric_editor_parsing_does_not_echo_input():
    assert parse_pairs("0,4\n4,7") == ((0, 4), (4, 7))
    assert parse_number("7") == 7
    with pytest.raises(SafetyError) as caught:
        parse_pairs("secret-value")
    assert "secret-value" not in str(caught.value)


def test_local_category_review_is_bounded_and_value_free_on_error(tmp_path):
    source = ROOT / "examples/synthetic_students.xlsx"
    assert local_categories(source, "campus") == ("Mars", "Moon")
    unsafe = tmp_path / "unsafe.xlsx"
    unsafe.write_text("field\nsynthetic@example.invalid\n")
    with pytest.raises(SafetyError) as caught:
        local_categories(unsafe, "field")
    assert "synthetic@example.invalid" not in str(caught.value)


def test_local_category_review_reports_blanks_without_returning_them(tmp_path):
    source = tmp_path / "synthetic-blanks.xlsx"
    source.write_bytes(
        excel_bytes(
            Table(
                ("synthetic_key", "category"),
                (
                    {"synthetic_key": "SYNTH-001", "category": "Alpha"},
                    {"synthetic_key": "SYNTH-002", "category": ""},
                    {"synthetic_key": "SYNTH-003", "category": "   "},
                ),
            )
        )
    )
    review = local_category_review(source, "category")
    assert review.values == ("Alpha",)
    assert review.blank_count == 2
    assert "" not in review.values and "   " not in review.values
    with pytest.raises(SafetyError):
        local_categories(source, "category")
