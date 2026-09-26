"""Synthetic region discovery and explicit-selection safeguards."""

import socket

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table as ExcelTable

from safeset.desktop_bridge import Bridge
from safeset.desktop_flow import (
    approve_relational_protection,
    approve_relational_reconstruction,
    prepare_relational_protection,
    prepare_relational_reconstruction,
)
from safeset.errors import SafetyError
from safeset.ingestion import Table, excel_workbook_bytes, read_excel_sheets
from safeset.policy_authoring import RuleDraft
from safeset.regions import discover_regions, read_confirmed_regions

from .conftest import PASSPHRASE
from .test_desktop_bridge import call


def _workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Research"
    sheet["A1"] = "Synthetic study title"
    sheet.append([])
    for row in (
        ("Person", "Group"),
        ("Synthetic A", "North"),
        ("Synthetic B", "South"),
    ):
        sheet.append(row)
    sheet["A8"] = "Notes outside selected data"
    sheet["D3"] = "Person"
    sheet["E3"] = "Score"
    sheet["D4"] = "Synthetic A"
    sheet["E4"] = "5"
    sheet["D5"] = "Synthetic B"
    sheet["E5"] = "6"
    book.save(path)


def test_confirmed_regions_have_independent_schemas_and_exclude_notes(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    _workbook(path)
    choices = {
        "Groups": {"sheet": "Research", "range": "A3:B5"},
        "Scores": {"sheet": "Research", "range": "D3:E5"},
    }
    tables = read_confirmed_regions(path, choices)
    assert tables["Groups"].columns == ("Person", "Group")
    assert tables["Scores"].columns == ("Person", "Score")
    assert len(tables["Groups"].rows) == 2
    assert "Notes outside selected data" not in str(tables)


def test_region_selection_rejects_overlap_and_ambiguous_heading(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    _workbook(path)
    with pytest.raises(SafetyError, match="overlap"):
        read_confirmed_regions(path, {
            "First": {"sheet": "Research", "range": "A3:B5"},
            "Second": {"sheet": "Research", "range": "B3:E5"},
        })
    with pytest.raises(SafetyError, match="headings"):
        read_confirmed_regions(path, {
            "Wrong": {"sheet": "Research", "range": "A1:B5"},
        })


def test_discovery_is_advisory(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    _workbook(path)
    proposals = discover_regions(path)
    assert all(set(item) == {"sheet", "range", "header_row", "kind"} for item in proposals)
    assert all("Synthetic A" not in str(item) for item in proposals)


def test_formal_excel_table_must_be_selected_whole(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    _workbook(path)
    book = load_workbook(path)
    book["Research"].add_table(ExcelTable(displayName="SyntheticGroups", ref="A3:B5"))
    book.save(path)
    proposals = discover_regions(path)
    assert any(item["range"] == "A3:B5" and item["kind"] == "excel_table" for item in proposals)
    assert read_confirmed_regions(path, {
        "Groups": {"sheet": "Research", "range": "A3:B5"}
    })["Groups"].columns == ("Person", "Group")
    with pytest.raises(SafetyError, match="complete Excel Table"):
        read_confirmed_regions(path, {
            "Partial": {"sheet": "Research", "range": "A3:A5"}
        })


def test_region_bundle_round_trip_preserves_original_layout(tmp_path, monkeypatch):
    def denied(*_args, **_kwargs):
        raise AssertionError("Network access attempted")

    monkeypatch.setattr(socket, "socket", denied)
    monkeypatch.setattr(socket, "create_connection", denied)
    path = tmp_path / "synthetic.xlsx"
    _workbook(path)
    for name in ("maps", "exports", "private"):
        (tmp_path / name).mkdir(mode=0o700)
    regions = {
        "Groups": {"sheet": "Research", "range": "A3:B5"},
        "Scores": {"sheet": "Research", "range": "D3:E5"},
    }
    drafts = {
        "Groups": {
            "Person": RuleDraft("pseudonymise", "direct_identifier"),
            "Group": RuleDraft("keep", "analytical_attribute", ("North", "South")),
        },
        "Scores": {
            "Person": RuleDraft("pseudonymise", "direct_identifier"),
            "Score": RuleDraft("keep_numeric", "analytical_attribute", bounds=(0, 10)),
        },
    }
    protected = tmp_path / "exports/protected.xlsx"
    bundle = tmp_path / "maps/regions.enc"
    review = prepare_relational_protection(
        path, tuple(regions), drafts, "2", protected, bundle,
        "controlled_pseudonymisation", (),
        {"Groups": ["Group"], "Scores": ["Score"]}, regions,
    )
    approve_relational_protection(review, PASSPHRASE, approved=True)
    returned = tmp_path / "returned.xlsx"
    tables = read_excel_sheets(protected, tuple(regions))
    group_rows = [dict(row) for row in tables["Groups"].rows]
    group_rows[0]["Group"] = "South"
    tables["Groups"] = Table(tables["Groups"].columns, tuple(reversed(group_rows)))
    score_rows = [dict(row) for row in tables["Scores"].rows]
    score_rows[0]["Score"] = "7"
    tables["Scores"] = Table(tables["Scores"].columns, tuple(score_rows))
    returned.write_bytes(excel_workbook_bytes(tables))
    restored = tmp_path / "private/restored.xlsx"
    restored_review = prepare_relational_reconstruction(
        returned, path, bundle, restored, PASSPHRASE
    )
    assert restored_review.bundle["version"] == 5
    approve_relational_reconstruction(
        restored_review, restored_review.new_columns, authorised=True,
        approved_changes={"Groups": ("Group",), "Scores": ("Score",)},
    )
    book = load_workbook(restored)
    assert book["Research"]["B4"].value == "South"
    assert book["Research"]["E4"].value == "7"
    assert book["Research"]["A1"].value == "Synthetic study title"
    assert book["Research"]["A8"].value == "Notes outside selected data"
    original = load_workbook(path)
    original["Research"]["A1"] = "Changed synthetic title"
    original.save(path)
    with pytest.raises(SafetyError, match="does not match"):
        prepare_relational_reconstruction(
            returned, path, bundle, tmp_path / "private/second.xlsx", PASSPHRASE
        )


def test_desktop_region_confirmation_binds_source_file(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    _workbook(path)
    bridge = Bridge()
    discovery = call(bridge, "discover_regions", {"path": str(path)})["result"]
    assert len(discovery["source_digest"]) == 64
    region = {"sheet": "Research", "range": "A3:B5"}
    inspection = call(bridge, "inspect_region", {
        "source": str(path), "name": "Groups", "region": region,
    })["result"]
    assert len(inspection["columns"]) == 2
    response = call(bridge, "prepare_relational_protection", {
        "source": str(path), "sheets": ["Groups"],
        "output": str(tmp_path / "protected.xlsx"),
        "drafts": {"Groups": {
            "Person": {"action": "pseudonymise", "classification": "direct_identifier",
                       "allowed_values": [], "bins": [], "bounds": None},
            "Group": {"action": "keep", "classification": "analytical_attribute",
                      "allowed_values": ["North", "South"], "bins": [], "bounds": None},
        }},
        "threshold": "2", "validation_profile": "controlled_pseudonymisation",
        "editable_fields": {"Groups": ["Group"]},
        "regions": {"Groups": region},
        "region_digest": "0" * 64,
    })
    assert not response["ok"]
    assert response["error"] == "region_review_stale"
