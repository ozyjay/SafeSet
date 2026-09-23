"""Synthetic version 2 working-copy round trip and failure cases."""

import socket
from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from safeset.bundle import decrypt_bundle, encrypt_bundle
from safeset.cli import app
from safeset.desktop_flow import (
    approve_protection,
    approve_reconstruction,
    prepare_protection,
    prepare_reconstruction,
)
from safeset.errors import SafetyError
from safeset.ingestion import (
    Table,
    excel_bytes,
    excel_workbook_bytes,
    read_excel,
    read_excel_sheets,
)
from safeset.policy_authoring import load_drafts
from safeset.pseudonyms import new_id
from safeset.reconstruction import reconstruct

from .conftest import PASSPHRASE, ROOT


def _set_formula_cache(
    path, worksheet_entry: str, coordinate: str, result: str | None
) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    original = BytesIO(path.read_bytes())
    updated = BytesIO()
    with ZipFile(original) as source, ZipFile(updated, "w") as destination:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == worksheet_entry:
                tree = ElementTree.fromstring(data)
                cell = tree.find(f".//{{{namespace}}}c[@r='{coordinate}']")
                assert cell is not None
                value = cell.find(f"{{{namespace}}}v")
                assert value is not None
                if result is None:
                    cell.set("t", "str")
                    cell.remove(value)
                else:
                    value.text = result
                data = ElementTree.tostring(tree, encoding="utf-8")
            destination.writestr(entry, data)
    path.write_bytes(updated.getvalue())


def _prepared(destinations):
    source = ROOT / "examples/synthetic_students.xlsx"
    drafts, threshold = load_drafts(source, ROOT / "examples/example-policy.yaml")
    output, bundle = destinations
    review = prepare_protection(source, drafts, str(threshold), output, bundle)
    assert review.validation.passed
    approve_protection(review, PASSPHRASE, approved=True)
    return source, output, bundle


def test_reconstruction_round_trip(destinations):
    source, output, bundle_path = _prepared(destinations)
    protected = read_excel(output)
    assert "student_name" not in protected.columns
    assert "student_number" not in protected.columns
    rows = tuple({**row, "Team": "Robot Team"} for row in protected.rows)
    returned_path = output.parent / "analysed.xlsx"
    returned_path.write_bytes(excel_bytes(Table((*protected.columns, "Team"), rows)))
    restored_path = output.parent.parent / "private/restored.xlsx"
    review = prepare_reconstruction(returned_path, source, bundle_path, restored_path, PASSPHRASE)
    assert review.new_columns == ("Team",)
    with pytest.raises(SafetyError):
        approve_reconstruction(review, (), authorised=True)
    assert approve_reconstruction(review, ("Team",), authorised=True) == 4
    original = read_excel(source)
    restored = read_excel(restored_path)
    assert restored.columns == (*original.columns, "Team")
    assert all(row["Team"] == "Robot Team" for row in restored.rows)
    assert [tuple(row[c] for c in original.columns) for row in restored.rows] == [
        tuple(row[c] for c in original.columns) for row in original.rows
    ]
    with pytest.raises(SafetyError):
        approve_reconstruction(review, ("Team",), authorised=True)


def test_reconstruction_combines_selected_worksheets_in_workbook_order(
    destinations, tmp_path
):
    original = read_excel(ROOT / "examples/synthetic_students.xlsx")
    source = tmp_path / "synthetic-multi-source.xlsx"
    source.write_bytes(
        excel_workbook_bytes(
            {
                "Earlier": Table(original.columns, original.rows[:2]),
                "Later": Table(original.columns, original.rows[2:]),
            }
        )
    )
    selected_source = ("Earlier", "Later")
    drafts, threshold = load_drafts(
        source, ROOT / "examples/example-policy.yaml", selected_source
    )
    output, bundle_path = destinations
    protection = prepare_protection(
        source,
        drafts,
        str(threshold),
        output,
        bundle_path,
        sheet=selected_source,
    )
    assert protection.validation.passed
    approve_protection(protection, PASSPHRASE, approved=True)

    protected = read_excel(output)
    result_columns = (*protected.columns, "Team")
    returned = output.parent / "synthetic-multi-returned.xlsx"
    returned.write_bytes(
        excel_workbook_bytes(
            {
                "First results": Table(
                    result_columns,
                    tuple({**row, "Team": "Robot Team"} for row in protected.rows[:2]),
                ),
                "Second results": Table(
                    result_columns,
                    tuple({**row, "Team": "Robot Team"} for row in protected.rows[2:]),
                ),
            }
        )
    )
    restored = output.parent.parent / "private/restored-multi-source.xlsx"
    review = prepare_reconstruction(
        returned,
        source,
        bundle_path,
        restored,
        PASSPHRASE,
        returned_sheet=("First results", "Second results"),
        source_sheet=selected_source,
    )
    assert review.new_columns == ("Team",)
    assert approve_reconstruction(review, ("Team",), authorised=True) == 4
    restored_table = read_excel(restored)
    assert [row["student_number"] for row in restored_table.rows] == [
        row["student_number"] for row in original.rows
    ]
    assert all(row["Team"] == "Robot Team" for row in restored_table.rows)


def test_reconstruction_copies_explicitly_approved_analysis_worksheet(destinations):
    source, output, bundle_path = _prepared(destinations)
    protected = read_excel(output)
    returned = output.parent / "analysed-with-sheet.xlsx"
    returned.write_bytes(
        excel_workbook_bytes(
            {
                "SafeSet": protected,
                "Changes": Table(
                    ("Finding", "Count"),
                    (
                        {"Finding": "Campus mismatch", "Count": "1"},
                        {"Finding": "No change", "Count": ""},
                    ),
                ),
            }
        )
    )
    restored = output.parent.parent / "private/restored-with-sheet.xlsx"
    review = prepare_reconstruction(
        returned,
        source,
        bundle_path,
        restored,
        PASSPHRASE,
        returned_sheet="SafeSet",
    )
    assert tuple(review.analysis_sheets) == ("Changes",)
    with pytest.raises(SafetyError, match="explicit approval"):
        approve_reconstruction(review, (), authorised=True)
    assert not restored.exists()
    assert approve_reconstruction(review, (), ("Changes",), authorised=True) == 4
    tables = read_excel_sheets(restored, ("SafeSet", "Changes"))
    assert tables["Changes"].rows == (
        {"Finding": "Campus mismatch", "Count": "1"},
        {"Finding": "No change", "Count": ""},
    )


def test_reconstruction_copies_saved_analysis_formula_result_as_static_value(destinations):
    source, output, bundle_path = _prepared(destinations)
    returned = output.parent / "analysed-formula-sheet.xlsx"
    workbook = load_workbook(output)
    protected_name = workbook.active.title
    analysis = workbook.create_sheet("Calculated")
    analysis.append(("Finding", "Count"))
    analysis.append(("Synthetic total", "=1+1"))
    analysis.append(("Synthetic empty result", '=""'))
    workbook.save(returned)
    _set_formula_cache(returned, "xl/worksheets/sheet2.xml", "B2", "2")
    _set_formula_cache(returned, "xl/worksheets/sheet2.xml", "B3", None)

    restored = output.parent.parent / "private/restored-formula-sheet.xlsx"
    review = prepare_reconstruction(
        returned,
        source,
        bundle_path,
        restored,
        PASSPHRASE,
        returned_sheet=protected_name,
    )
    assert review.analysis_sheets["Calculated"].formula_cells == 2
    assert review.analysis_sheets["Calculated"].rows == (
        {"Finding": "Synthetic total", "Count": "2"},
        {"Finding": "Synthetic empty result", "Count": ""},
    )

    assert approve_reconstruction(review, (), ("Calculated",), authorised=True) == 4
    restored_analysis = read_excel_sheets(restored, ("Calculated",))["Calculated"]
    assert restored_analysis.formula_cells == 0
    assert restored_analysis.rows == (
        {"Finding": "Synthetic total", "Count": "2"},
        {"Finding": "Synthetic empty result", "Count": ""},
    )


def test_reconstruction_rejects_formula_in_bound_returned_sheet(destinations):
    source, output, bundle_path = _prepared(destinations)
    returned = output.parent / "formula-in-protected-sheet.xlsx"
    workbook = load_workbook(output)
    worksheet = workbook.active
    result_column = worksheet.max_column + 1
    worksheet.cell(1, result_column, "Team")
    for row in range(2, worksheet.max_row + 1):
        worksheet.cell(row, result_column, '=CONCAT("Robot", " Team")')
    workbook.save(returned)

    with pytest.raises(SafetyError, match="unsupported cell type"):
        prepare_reconstruction(
            returned,
            source,
            bundle_path,
            output.parent.parent / "private/rejected-formula.xlsx",
            PASSPHRASE,
        )


def test_reconstruction_rejects_analysis_formula_without_saved_result(destinations):
    source, output, bundle_path = _prepared(destinations)
    returned = output.parent / "uncalculated-formula-sheet.xlsx"
    workbook = load_workbook(output)
    protected_name = workbook.active.title
    analysis = workbook.create_sheet("Calculated")
    analysis.append(("Finding", "Count"))
    analysis.append(("Synthetic total", "=1+1"))
    workbook.save(returned)

    with pytest.raises(SafetyError, match="formula has no saved result"):
        prepare_reconstruction(
            returned,
            source,
            bundle_path,
            output.parent.parent / "private/rejected-uncalculated.xlsx",
            PASSPHRASE,
            returned_sheet=protected_name,
        )


def test_analysis_worksheet_change_after_review_blocks_publication(destinations):
    source, output, bundle_path = _prepared(destinations)
    protected = read_excel(output)
    returned = output.parent / "reviewed-analysis.xlsx"
    returned.write_bytes(
        excel_workbook_bytes(
            {
                "SafeSet": protected,
                "Changes": Table(
                    ("Finding", "Count"),
                    ({"Finding": "Campus mismatch", "Count": "1"},),
                ),
            }
        )
    )
    restored = output.parent.parent / "private/stale-analysis.xlsx"
    review = prepare_reconstruction(
        returned,
        source,
        bundle_path,
        restored,
        PASSPHRASE,
        returned_sheet="SafeSet",
    )
    returned.write_bytes(
        excel_workbook_bytes(
            {
                "SafeSet": protected,
                "Changes": Table(
                    ("Finding", "Count"),
                    ({"Finding": "Campus mismatch", "Count": "2"},),
                ),
            }
        )
    )
    with pytest.raises(SafetyError, match="changed after restoration review"):
        approve_reconstruction(review, (), ("Changes",), authorised=True)
    assert not restored.exists()


@pytest.mark.parametrize(
    "change", ["duplicate", "missing", "unknown", "malformed", "changed", "extra"]
)
def test_reconstruction_rejects_tampering(destinations, change):
    source, output, bundle_path = _prepared(destinations)
    from safeset.bundle import read_bundle

    bundle = read_bundle(bundle_path, PASSPHRASE)
    protected = read_excel(output)
    rows = [dict(row) for row in protected.rows]
    columns = protected.columns
    if change == "duplicate":
        rows[0]["record_id"] = rows[1]["record_id"]
    elif change == "missing":
        rows.pop()
    elif change == "unknown":
        rows[0]["record_id"] = new_id()
    elif change == "malformed":
        rows[0]["record_id"] = "bad-id"
    elif change == "changed":
        rows[0]["campus"] = new_id()
    else:
        columns = (*columns, "student_name")
        rows = [{**row, "student_name": "Invented Person"} for row in rows]
    with pytest.raises(SafetyError):
        reconstruct(read_excel(source), Table(columns, tuple(rows)), bundle, ())


def test_wrong_source_and_legacy_map_rejected(destinations):
    source, output, bundle_path = _prepared(destinations)
    from safeset.bundle import read_bundle

    bundle = read_bundle(bundle_path, PASSPHRASE)
    original = read_excel(source)
    rows = deepcopy(original.rows)
    rows[0]["student_name"] = "Invented Changed"
    with pytest.raises(SafetyError, match="source does not match"):
        reconstruct(Table(original.columns, rows), read_excel(output), bundle, ())
    with pytest.raises(SafetyError):
        decrypt_bundle(b"SAFESET1\n" + bundle_path.read_bytes()[9:], PASSPHRASE)
    ciphertext = encrypt_bundle(bundle, PASSPHRASE)
    assert decrypt_bundle(ciphertext, PASSPHRASE) == bundle
    with pytest.raises(SafetyError):
        decrypt_bundle(ciphertext[:-1] + b"!", PASSPHRASE)


def test_bundle_is_minimal_and_runtime_offline(destinations, monkeypatch):
    def denied(*_args, **_kwargs):
        raise AssertionError("Runtime attempted a network connection")

    monkeypatch.setattr(socket, "socket", denied)
    monkeypatch.setattr(socket, "create_connection", denied)
    source, output, bundle_path = _prepared(destinations)
    from safeset.bundle import read_bundle

    bundle = read_bundle(bundle_path, PASSPHRASE)
    assert set(bundle) == {
        "version",
        "export_id",
        "source_digest",
        "source_column",
        "source_columns",
        "protected_columns",
        "records",
        "codebooks",
        "policy",
    }
    assert set(bundle["codebooks"]) == {"campus", "subject"}
    encoded = str(bundle)
    for row in read_excel(source).rows:
        assert row["student_name"] not in encoded
        assert row["email"] not in encoded
        assert row["notes"] not in encoded
        assert row["student_number"].encode() not in bundle_path.read_bytes()
    assert "student_number" not in read_excel(output).columns


def test_source_change_after_review_blocks_publication(destinations, tmp_path):
    source = tmp_path / "source.xlsx"
    original = read_excel(ROOT / "examples/synthetic_students.xlsx")
    source.write_bytes(excel_bytes(original))
    drafts, threshold = load_drafts(source, ROOT / "examples/example-policy.yaml")
    output, bundle = destinations
    review = prepare_protection(source, drafts, str(threshold), output, bundle)
    rows = deepcopy(original.rows)
    rows[0]["student_name"] = "Invented Change"
    source.write_bytes(excel_bytes(Table(original.columns, rows)))
    with pytest.raises(SafetyError, match="changed"):
        approve_protection(review, PASSPHRASE, approved=True)
    assert not output.exists() and not bundle.exists()


def test_cli_protect_and_reconstruct(destinations, monkeypatch):
    monkeypatch.setattr("safeset.cli.secret", lambda **_kwargs: PASSPHRASE)
    output, bundle_path = destinations
    source = ROOT / "examples/synthetic_students.xlsx"
    runner = CliRunner()
    protected = runner.invoke(
        app,
        [
            "protect",
            str(source),
            "--policy",
            str(ROOT / "examples/example-policy.yaml"),
            "--output",
            str(output),
            "--create-bundle",
            "--bundle",
            str(bundle_path),
            "--approve-export",
        ],
    )
    assert protected.exit_code == 0, protected.output
    table = read_excel(output)
    returned = output.parent / "with-results.xlsx"
    returned.write_bytes(
        excel_workbook_bytes(
            {
                "SafeSet": Table(
                (*table.columns, "Team"),
                tuple({**row, "Team": "Robot Team"} for row in table.rows),
                ),
                "Summary": Table(
                    ("Finding", "Count"),
                    ({"Finding": "No change", "Count": "4"},),
                ),
            }
        )
    )
    restored = output.parent.parent / "private/reconstructed.xlsx"
    missing_approval = runner.invoke(
        app,
        [
            "reconstruct",
            str(returned),
            "--original-source",
            str(source),
            "--bundle",
            str(bundle_path),
            "--output",
            str(restored),
            "--sheet",
            "SafeSet",
            "--authorise",
        ],
    )
    assert missing_approval.exit_code != 0
    assert not restored.exists()
    result = runner.invoke(
        app,
        [
            "reconstruct",
            str(returned),
            "--original-source",
            str(source),
            "--bundle",
            str(bundle_path),
            "--output",
            str(restored),
            "--result-column",
            "Team",
            "--sheet",
            "SafeSet",
            "--analysis-sheet",
            "Summary",
            "--authorise",
        ],
    )
    assert result.exit_code == 0, result.output
    restored_tables = read_excel_sheets(restored, ("SafeSet", "Summary"))
    assert restored_tables["SafeSet"].columns == (*read_excel(source).columns, "Team")
    assert restored_tables["Summary"].rows == ({"Finding": "No change", "Count": "4"},)
