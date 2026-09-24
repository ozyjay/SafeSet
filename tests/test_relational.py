"""Synthetic multi-sheet linkage, validation and reconstruction tests."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from safeset.desktop_bridge import Bridge
from safeset.desktop_flow import (
    approve_relational_protection,
    approve_relational_reconstruction,
    locate_unsafe_source_cells,
    prepare_relational_protection,
    prepare_relational_reconstruction,
)
from safeset.errors import SafetyError
from safeset.ingestion import Table, excel_workbook_bytes, read_excel_sheets
from safeset.policy import parse_policy
from safeset.policy_authoring import RuleDraft
from safeset.relational import (
    create_relational_bundle,
    read_relational_bundle,
    sanitise_relational,
    validate_relational,
    validate_relational_bundle,
)

from .conftest import PASSPHRASE


def _sources() -> dict[str, Table]:
    return {
        "Enrolments": Table(
            ("student_key", "display_name", "campus", "cohort"),
            (
                {
                    "student_key": "SYNTH-001",
                    "display_name": "Invented Ada",
                    "campus": "North",
                    "cohort": "Alpha",
                },
                {
                    "student_key": "SYNTH-002",
                    "display_name": "Invented Beau",
                    "campus": "South",
                    "cohort": "Beta",
                },
            ),
        ),
        "Preferences": Table(
            ("person_ref", "stream", "cohort"),
            (
                {"person_ref": "SYNTH-002", "stream": "Robotics", "cohort": "Beta"},
                {"person_ref": "SYNTH-001", "stream": "Design", "cohort": "Alpha"},
            ),
        ),
    }


def _policies():
    enrolments = {
        "version": 2,
        "min_group_size": 2,
        "columns": {
            "student_key": {
                "action": "pseudonymise",
                "classification": "direct_identifier",
            },
            "display_name": {"action": "drop", "classification": "direct_identifier"},
            "campus": {
                "action": "code",
                "classification": "quasi_identifier",
                "allowed_values": ["North", "South"],
            },
            "cohort": {
                "action": "code",
                "classification": "quasi_identifier",
                "allowed_values": ["Alpha", "Beta"],
            },
        },
    }
    return {
        "Enrolments": parse_policy(enrolments),
        "Preferences": parse_policy(
            {
                "version": 2,
                "min_group_size": 2,
                "columns": {
                    "person_ref": {
                        "action": "pseudonymise",
                        "classification": "direct_identifier",
                    },
                    "stream": {
                        "action": "code",
                        "classification": "quasi_identifier",
                        "allowed_values": ["Design", "Robotics"],
                    },
                    "cohort": {
                        "action": "code",
                        "classification": "quasi_identifier",
                        "allowed_values": ["Alpha", "Beta"],
                    },
                },
            }
        ),
    }


def _drafts():
    return {
        "Enrolments": {
            "student_key": RuleDraft("pseudonymise", "direct_identifier"),
            "display_name": RuleDraft("drop", "direct_identifier"),
            "campus": RuleDraft("code", "quasi_identifier", ("North", "South")),
            "cohort": RuleDraft("code", "quasi_identifier", ("Alpha", "Beta")),
        },
        "Preferences": {
            "person_ref": RuleDraft("pseudonymise", "direct_identifier"),
            "stream": RuleDraft("code", "quasi_identifier", ("Design", "Robotics")),
            "cohort": RuleDraft("code", "quasi_identifier", ("Alpha", "Beta")),
        },
    }


def _write_source(path: Path) -> None:
    path.write_bytes(excel_workbook_bytes(_sources()))


def _draft_payload() -> dict:
    return {
        sheet: {
            name: {
                "action": draft.action,
                "classification": draft.classification,
                "allowed_values": list(draft.allowed_values),
                "bins": [list(pair) for pair in draft.bins],
                "bounds": list(draft.bounds) if draft.bounds is not None else None,
            }
            for name, draft in fields.items()
        }
        for sheet, fields in _drafts().items()
    }


def _call(bridge: Bridge, command: str, payload: dict) -> dict:
    line = (
        json.dumps(
            {"version": 1, "id": "relational-test", "command": command, "payload": payload}
        ).encode()
        + b"\n"
    )
    return json.loads(bridge.process_line(line))


def test_shared_entities_have_fresh_row_ids_and_controlled_warnings():
    candidate = sanitise_relational(_sources(), _policies())
    enrolments = candidate.tables["Enrolments"].rows
    preferences = candidate.tables["Preferences"].rows
    assert enrolments[0]["entity_id"] == preferences[1]["entity_id"]
    assert enrolments[1]["entity_id"] == preferences[0]["entity_id"]
    assert {row["record_id"] for table in candidate.tables.values() for row in table.rows}
    assert enrolments[0]["record_id"] != preferences[1]["record_id"]
    strict = validate_relational(candidate, _policies(), "strict")
    assert not strict.passed and strict.linked_small_entities == 2
    controlled = validate_relational(candidate, _policies(), "controlled_pseudonymisation")
    assert controlled.passed
    assert controlled.linked_small_entities == 2
    assert controlled.linked_warnings


def test_relational_exact_numeric_field_preserves_blanks():
    sources = {
        "Marks": Table(
            ("student_key", "score"),
            (
                {"student_key": "SYNTH-001", "score": ""},
                {"student_key": "SYNTH-002", "score": ""},
                {"student_key": "SYNTH-003", "score": "4.20"},
                {"student_key": "SYNTH-004", "score": "4.20"},
            ),
        )
    }
    policies = {
        "Marks": parse_policy(
            {
                "version": 3,
                "min_group_size": 2,
                "columns": {
                    "student_key": {
                        "action": "pseudonymise",
                        "classification": "direct_identifier",
                    },
                    "score": {
                        "action": "keep_numeric",
                        "classification": "analytical_attribute",
                        "bounds": [0, 10],
                    },
                },
            }
        )
    }
    candidate = sanitise_relational(sources, policies)
    assert [row["score"] for row in candidate.tables["Marks"].rows] == [
        "",
        "",
        "4.2",
        "4.2",
    ]
    assert validate_relational(candidate, policies, "controlled_pseudonymisation").passed


def test_shared_obfuscation_requires_explicit_field_confirmation():
    independent = sanitise_relational(_sources(), _policies())
    assert (
        independent.codebooks["Enrolments"]["cohort"]["Alpha"]
        != independent.codebooks["Preferences"]["cohort"]["Alpha"]
    )

    shared = sanitise_relational(_sources(), _policies(), ("cohort",))
    assert shared.shared_code_fields == ("cohort",)
    assert (
        shared.codebooks["Enrolments"]["cohort"]["Alpha"]
        == shared.codebooks["Preferences"]["cohort"]["Alpha"]
    )
    validation = validate_relational(shared, _policies(), "controlled_pseudonymisation")
    assert any("Shared obfuscation codebooks" in warning for warning in validation.linked_warnings)


def test_shared_obfuscation_bundle_validation_and_legacy_compatibility():
    shared = sanitise_relational(_sources(), _policies(), ("cohort",))
    bundle = create_relational_bundle(
        _sources(), _policies(), shared, "controlled_pseudonymisation"
    )
    assert validate_relational_bundle(bundle)["shared_code_fields"] == ["cohort"]
    bundle["sheets"]["Preferences"]["codebooks"]["cohort"]["Alpha"] = (
        "00000000-0000-4000-8000-000000000001"
    )
    with pytest.raises(SafetyError, match="shared category codebook"):
        validate_relational_bundle(bundle)

    independent = sanitise_relational(_sources(), _policies())
    legacy = create_relational_bundle(
        _sources(), _policies(), independent, "controlled_pseudonymisation"
    )
    legacy.pop("shared_code_fields")
    assert validate_relational_bundle(legacy)["shared_code_fields"] == []


def test_shared_obfuscation_rejects_unconfirmed_or_ineligible_fields():
    with pytest.raises(SafetyError, match="shared obfuscation field"):
        sanitise_relational(_sources(), _policies(), ("campus",))


def test_relational_source_text_locator_reports_only_bound_cell_coordinates(tmp_path):
    tables = _sources()
    rows = [dict(row) for row in tables["Enrolments"].rows]
    rows[0]["display_name"] = "SYNTHETIC\tUNSAFE"
    tables["Enrolments"] = Table(tables["Enrolments"].columns, tuple(rows))
    source = tmp_path / "synthetic-related.xlsx"
    source.write_bytes(excel_workbook_bytes(tables))
    for name in ("exports", "maps"):
        (tmp_path / name).mkdir(mode=0o700)
    protected = tmp_path / "exports/protected.xlsx"
    bundle_path = tmp_path / "maps/related.enc"
    review = prepare_relational_protection(
        source, ("Enrolments", "Preferences"), _drafts(), "2", protected,
        bundle_path, "controlled_pseudonymisation",
    )
    assert review.validation.passed
    approve_relational_protection(review, PASSPHRASE, approved=True)
    located = locate_unsafe_source_cells(
        source, bundle_path, PASSPHRASE, relational=True
    )
    assert located == {"count": 1, "cells": [{"sheet": "Enrolments", "cell": "B2"}]}
    assert "SYNTHETIC" not in str(located)
    private = tmp_path / "private"
    private.mkdir(mode=0o700)
    with pytest.raises(SafetyError, match="Original source contains unsafe spreadsheet text"):
        prepare_relational_reconstruction(
            protected, source, bundle_path, private / "restored.xlsx", PASSPHRASE
        )


def test_relational_round_trip_and_minimal_bundle(tmp_path):
    source = tmp_path / "synthetic-related.xlsx"
    _write_source(source)
    exports = tmp_path / "exports"
    maps = tmp_path / "maps"
    private = tmp_path / "private"
    for directory in (exports, maps, private):
        directory.mkdir(mode=0o700)
    protected = exports / "protected.xlsx"
    bundle_path = maps / "related.enc"
    review = prepare_relational_protection(
        source,
        ("Enrolments", "Preferences"),
        _drafts(),
        "2",
        protected,
        bundle_path,
        "controlled_pseudonymisation",
        ("cohort",),
    )
    assert review.validation.passed
    approve_relational_protection(review, PASSPHRASE, approved=True)
    protected_tables = read_excel_sheets(protected, ("Enrolments", "Preferences"))
    assert (
        protected_tables["Enrolments"].rows[0]["entity_id"]
        == (protected_tables["Preferences"].rows[1]["entity_id"])
    )
    bundle = read_relational_bundle(bundle_path, PASSPHRASE)
    assert bundle["version"] == 3
    assert bundle["shared_code_fields"] == ["cohort"]
    assert set(bundle["entities"].values()) == {"SYNTH-001", "SYNTH-002"}
    assert "North" not in str(bundle["entities"])
    assert "Invented Ada" not in str(bundle)
    returned = exports / "returned.xlsx"
    returned.write_bytes(
        excel_workbook_bytes(
            {
                name: Table(
                    (*table.columns, "Team"),
                    tuple({**row, "Team": "Koala Team"} for row in table.rows),
                )
                for name, table in protected_tables.items()
            }
        )
    )
    restored = private / "restored.xlsx"
    restore_review = prepare_relational_reconstruction(
        returned, source, bundle_path, restored, PASSPHRASE
    )
    assert restore_review.new_columns == {
        "Enrolments": ("Team",),
        "Preferences": ("Team",),
    }
    assert (
        approve_relational_reconstruction(
            restore_review,
            {"Enrolments": ("Team",), "Preferences": ("Team",)},
            authorised=True,
        )
        == 4
    )
    restored_tables = read_excel_sheets(restored, ("Enrolments", "Preferences"))
    assert restored_tables["Enrolments"].columns == (
        "student_key",
        "display_name",
        "campus",
        "cohort",
        "Team",
    )
    assert restored_tables["Preferences"].columns == (
        "person_ref",
        "stream",
        "cohort",
        "Team",
    )


def test_relational_reconstruction_rejects_changed_entity_link(tmp_path):
    source = tmp_path / "synthetic-related.xlsx"
    _write_source(source)
    for name in ("exports", "maps", "private"):
        (tmp_path / name).mkdir(mode=0o700)
    protected = tmp_path / "exports/protected.xlsx"
    bundle_path = tmp_path / "maps/related.enc"
    review = prepare_relational_protection(
        source,
        ("Enrolments", "Preferences"),
        _drafts(),
        "2",
        protected,
        bundle_path,
        "controlled_pseudonymisation",
    )
    approve_relational_protection(review, PASSPHRASE, approved=True)
    tables = read_excel_sheets(protected, ("Enrolments", "Preferences"))
    rows = [dict(row) for row in tables["Preferences"].rows]
    rows[0]["entity_id"] = rows[1]["entity_id"]
    changed = tmp_path / "exports/changed.xlsx"
    changed.write_bytes(
        excel_workbook_bytes(
            {**tables, "Preferences": Table(tables["Preferences"].columns, tuple(rows))}
        )
    )
    with pytest.raises(SafetyError):
        prepare_relational_reconstruction(
            changed,
            source,
            bundle_path,
            tmp_path / "private/restored.xlsx",
            PASSPHRASE,
        )
    extra = tmp_path / "exports/extra-sheet.xlsx"
    extra.write_bytes(
        excel_workbook_bytes(
            {
                **tables,
                "Unapproved": Table(("note",), ({"note": "Synthetic only"},)),
            }
        )
    )
    restored = tmp_path / "private/extra-restored.xlsx"
    extra_review = prepare_relational_reconstruction(
        extra, source, bundle_path, restored, PASSPHRASE
    )
    assert tuple(extra_review.analysis_sheets) == ("Unapproved",)
    with pytest.raises(SafetyError, match="explicit approval"):
        approve_relational_reconstruction(
            extra_review,
            {"Enrolments": (), "Preferences": ()},
            authorised=True,
        )
    assert not restored.exists()
    assert (
        approve_relational_reconstruction(
            extra_review,
            {"Enrolments": (), "Preferences": ()},
            ("Unapproved",),
            authorised=True,
        )
        == 4
    )
    restored_tables = read_excel_sheets(
        restored, ("Enrolments", "Preferences", "Unapproved")
    )
    assert restored_tables["Unapproved"].rows == ({"note": "Synthetic only"},)


def test_relational_reconstruction_rejects_unsafe_added_worksheet(tmp_path):
    source = tmp_path / "synthetic-related.xlsx"
    _write_source(source)
    for name in ("exports", "maps", "private"):
        (tmp_path / name).mkdir(mode=0o700)
    protected = tmp_path / "exports/protected.xlsx"
    bundle_path = tmp_path / "maps/related.enc"
    review = prepare_relational_protection(
        source,
        ("Enrolments", "Preferences"),
        _drafts(),
        "2",
        protected,
        bundle_path,
        "controlled_pseudonymisation",
    )
    approve_relational_protection(review, PASSPHRASE, approved=True)
    tables = read_excel_sheets(protected, ("Enrolments", "Preferences"))
    returned = tmp_path / "exports/unsafe-analysis.xlsx"
    returned.write_bytes(
        excel_workbook_bytes(
            {
                **tables,
                "Analysis": Table(
                    ("note",),
                    ({"note": "This synthetic analysis text is deliberately too long"},),
                ),
            }
        )
    )
    with pytest.raises(SafetyError, match="unsafe content"):
        prepare_relational_reconstruction(
            returned,
            source,
            bundle_path,
            tmp_path / "private/restored.xlsx",
            PASSPHRASE,
        )


def test_relational_reconstruction_uses_bundle_sheets_and_rejects_missing_or_hidden(tmp_path):
    source = tmp_path / "synthetic-related.xlsx"
    _write_source(source)
    for name in ("exports", "maps", "private"):
        (tmp_path / name).mkdir(mode=0o700)
    protected = tmp_path / "exports/protected.xlsx"
    bundle_path = tmp_path / "maps/related.enc"
    review = prepare_relational_protection(
        source,
        ("Enrolments", "Preferences"),
        _drafts(),
        "2",
        protected,
        bundle_path,
        "controlled_pseudonymisation",
    )
    approve_relational_protection(review, PASSPHRASE, approved=True)
    tables = read_excel_sheets(protected, ("Enrolments", "Preferences"))

    automatic = prepare_relational_reconstruction(
        protected,
        source,
        bundle_path,
        tmp_path / "private/automatic.xlsx",
        PASSPHRASE,
    )
    assert tuple(automatic.returned_tables) == ("Enrolments", "Preferences")
    assert automatic.analysis_sheets == {}

    missing = tmp_path / "exports/missing-required.xlsx"
    missing.write_bytes(excel_workbook_bytes({"Enrolments": tables["Enrolments"]}))
    with pytest.raises(SafetyError, match="coverage does not match"):
        prepare_relational_reconstruction(
            missing,
            source,
            bundle_path,
            tmp_path / "private/missing-bound-sheet.xlsx",
            PASSPHRASE,
        )

    hidden = tmp_path / "exports/hidden-required.xlsx"
    workbook = load_workbook(protected)
    workbook["Preferences"].sheet_state = "hidden"
    workbook.save(hidden)
    with pytest.raises(SafetyError, match="hidden worksheets"):
        prepare_relational_reconstruction(
            hidden,
            source,
            bundle_path,
            tmp_path / "private/hidden-bound-sheet.xlsx",
            PASSPHRASE,
        )


def test_relational_desktop_bridge_review_and_publication(tmp_path):
    source = tmp_path / "synthetic-related.xlsx"
    _write_source(source)
    for name in ("exports", "maps"):
        (tmp_path / name).mkdir(mode=0o700)
    output = tmp_path / "exports/protected.xlsx"
    bundle = tmp_path / "maps/related.enc"
    bridge = Bridge()
    prepared = _call(
        bridge,
        "prepare_relational_protection",
        {
            "source": str(source),
            "sheets": ["Enrolments", "Preferences"],
            "output": str(output),
            "bundle": str(bundle),
            "drafts": _draft_payload(),
            "threshold": "2",
            "validation_profile": "controlled_pseudonymisation",
            "shared_code_fields": ["cohort"],
        },
    )
    assert prepared["ok"]
    result = prepared["result"]
    assert result["worksheets"] == 2 and result["entities"] == 2
    assert result["shared_code_fields"] == ["cohort"]
    assert result["validation"]["passed"]
    assert "SYNTH-001" not in json.dumps(prepared)
    approved = _call(
        bridge,
        "approve_relational_protection",
        {"review_id": result["review_id"], "passphrase": PASSPHRASE},
    )
    assert approved["ok"] and output.exists() and bundle.exists()
