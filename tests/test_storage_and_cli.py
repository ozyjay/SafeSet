import getpass
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from safeset.cli import app, secret
from safeset.errors import SafetyError
from safeset.ingestion import Table, excel_bytes, read_excel
from safeset.mapping import read_mapping
from safeset.storage import (
    default_map_path,
    map_destination,
    output_destination,
    private_directory,
    publish,
)
from safeset.workflow import export_candidate

from .conftest import PASSPHRASE, ROOT


def test_default_map_outside_checkout_and_not_export():
    path = default_map_path()
    assert not path.is_relative_to(ROOT)
    assert path.parent == Path.home() / ".local/share/safeset/maps"
    assert default_map_path() != path


def test_cli_inspect_requires_sheet_for_multi_sheet_workbook(tmp_path):
    path = tmp_path / "multiple.xlsx"
    workbook = load_workbook(ROOT / "examples/synthetic_students.xlsx")
    workbook.active.title = "Allocations"
    workbook.create_sheet("Instructions").append(("Invented help",))
    workbook.save(path)
    runner = CliRunner()
    missing = runner.invoke(app, ["inspect", str(path)])
    assert missing.exit_code != 0
    assert "Select one or more worksheets" in missing.output
    selected = runner.invoke(app, ["inspect", str(path), "--sheet", "Allocations"])
    assert selected.exit_code == 0
    assert '"rows": 4' in selected.output


def test_cli_accepts_repeated_sheet_options(tmp_path):
    path = tmp_path / "multiple.xlsx"
    workbook = load_workbook(ROOT / "examples/synthetic_students.xlsx")
    workbook.active.title = "Earlier"
    workbook.copy_worksheet(workbook.active).title = "Updated"
    workbook.save(path)
    result = CliRunner().invoke(
        app, ["inspect", str(path), "--sheet", "Earlier", "--sheet", "Updated"]
    )
    assert result.exit_code == 0
    assert '"rows": 8' in result.output


def test_repository_and_symlink_destination_rejected(tmp_path):
    with pytest.raises(SafetyError):
        output_destination(ROOT / "unsafe.xlsx")
    repo = tmp_path / "repo"
    repo.mkdir()
    (repo / ".git").mkdir()
    link = tmp_path / "link"
    link.symlink_to(repo, target_is_directory=True)
    with pytest.raises(SafetyError):
        map_destination(link / "map.enc", tmp_path / "out/safe.xlsx")
    with pytest.raises(SafetyError):
        output_destination(link / "out.xlsx")


def test_mapping_not_next_to_or_under_export(destinations):
    output, _ = destinations
    for path in [output.parent / "map.enc", output.parent / "sub/map.enc"]:
        with pytest.raises(SafetyError):
            map_destination(path, output)


def test_private_directory_and_map_permissions(candidate, policy, destinations):
    output, path = destinations
    path.parent.chmod(0o755)
    with pytest.raises(SafetyError):
        private_directory(path.parent, create=False)
    path.parent.chmod(0o700)
    export_candidate(
        candidate,
        policy,
        output,
        path,
        PASSPHRASE,
        approved=True,
        create_map=True,
        source_path=ROOT / "examples/synthetic_students.xlsx",
    )
    path.chmod(0o644)
    with pytest.raises(SafetyError):
        read_mapping(path, PASSPHRASE)


def test_no_overwrite_or_plaintext_staging(tmp_path):
    path = tmp_path / "present.xlsx"
    path.write_bytes(b"keep-existing")
    with pytest.raises(SafetyError):
        publish(path, b"replacement")
    assert path.read_bytes() == b"keep-existing"
    assert not list(tmp_path.glob(".safeset-*"))


def test_publication_failure_retains_only_encrypted_map(
    candidate, policy, destinations, monkeypatch
):
    output, path = destinations

    def fail_export(destination, data):
        if destination == output:
            raise SafetyError("simulated filesystem failure")
        publish(destination, data)

    monkeypatch.setattr("safeset.workflow.publish", fail_export)
    with pytest.raises(SafetyError, match="encrypted map was retained"):
        export_candidate(
            candidate,
            policy,
            output,
            path,
            PASSPHRASE,
            approved=True,
            create_map=True,
            source_path=ROOT / "examples/synthetic_students.xlsx",
        )
    assert not output.exists()
    assert read_mapping(path, PASSPHRASE) == candidate.mapping


@pytest.mark.parametrize("mode", ["missing_create", "declined", "bad_validation"])
def test_cli_no_output_without_approval_and_validity(destinations, tmp_path, monkeypatch, mode):
    def unexpected_secret(**kwargs):
        raise AssertionError("Secret prompted before approval/validation")

    monkeypatch.setattr("safeset.cli.secret", unexpected_secret)
    output, path = destinations
    source_path = ROOT / "examples/synthetic_students.xlsx"
    if mode == "bad_validation":
        source_path = tmp_path / "bad.xlsx"
        original = read_excel(ROOT / "examples/synthetic_students.xlsx")
        rows = tuple(
            {**row, "campus": "Unapproved"} if index == 0 else row
            for index, row in enumerate(original.rows)
        )
        source_path.write_bytes(excel_bytes(Table(original.columns, rows)))
    args = [
        "sanitise",
        str(source_path),
        "--policy",
        str(ROOT / "examples/example-policy.yaml"),
        "--output",
        str(output),
        "--map",
        str(path),
    ]
    if mode != "missing_create":
        args += ["--create-map"]
    result = CliRunner().invoke(app, args, input="n\n")
    assert result.exit_code != 0
    assert not isinstance(result.exception, AssertionError)
    assert not output.exists() and not path.exists()
    assert PASSPHRASE not in result.output and "Unapproved" not in result.output


def test_restore_requires_authorisation(destinations):
    output, path = destinations
    result = CliRunner().invoke(
        app,
        [
            "restore",
            str(output),
            "--map",
            str(path),
            "--output",
            str(output.parent / "restored.xlsx"),
        ],
    )
    assert result.exit_code != 0
    assert "--authorise" in result.output


def test_validate_failure_exit_status(candidate, tmp_path):
    candidate.table.rows[0]["record_id"] = "private-invalid-value"
    path = tmp_path / "bad.xlsx"
    path.write_bytes(excel_bytes(candidate.table))
    result = CliRunner().invoke(
        app, ["validate", str(path), "--policy", str(ROOT / "examples/example-policy.yaml")]
    )
    assert result.exit_code != 0
    assert "private-invalid-value" not in result.output


def test_secret_confirm_and_no_echo_fallback(monkeypatch):
    values = iter([PASSPHRASE, "different synthetic passphrase"])
    monkeypatch.setattr(getpass, "getpass", lambda prompt: next(values))
    with pytest.raises(SafetyError, match="do not match"):
        secret(confirm=True)

    def unavailable(prompt):
        import warnings

        warnings.warn("cannot hide input", getpass.GetPassWarning, stacklevel=2)

    monkeypatch.setattr(getpass, "getpass", unavailable)
    with pytest.raises(SafetyError, match="terminal"):
        secret()


def test_existing_map_prevents_export(candidate, policy, destinations):
    output, path = destinations
    path.write_bytes(b"existing synthetic sentinel")
    with pytest.raises(SafetyError):
        export_candidate(
            candidate,
            policy,
            output,
            path,
            PASSPHRASE,
            approved=True,
            create_map=True,
            source_path=ROOT / "examples/synthetic_students.xlsx",
        )
    assert not output.exists()
    assert path.read_bytes() == b"existing synthetic sentinel"


def test_expanded_export_limit_checked_before_writes(candidate, policy, destinations, monkeypatch):
    output, path = destinations
    monkeypatch.setattr("safeset.workflow.MAX_BYTES", 1)
    with pytest.raises(SafetyError, match="size limit"):
        export_candidate(
            candidate,
            policy,
            output,
            path,
            PASSPHRASE,
            approved=True,
            create_map=True,
            source_path=ROOT / "examples/synthetic_students.xlsx",
        )
    assert not path.exists() and not output.exists()
