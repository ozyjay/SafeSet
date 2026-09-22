from copy import deepcopy
from datetime import date, datetime, time
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table as ExcelTable

from safeset.errors import SafetyError
from safeset.ingestion import (
    MAX_FIELD,
    Table,
    excel_bytes,
    list_excel_sheets,
    read_bounded,
    read_excel,
)
from safeset.policy import load_policy, parse_policy
from safeset.transform import sanitise

from .conftest import ROOT


@pytest.mark.parametrize(
    "change",
    [
        lambda w: setattr(w.active["A1"], "value", "b"),
        lambda w: setattr(w.active["A1"], "value", " a"),
        lambda w: setattr(w.active["A2"], "value", "=" + "1+2"),
        lambda w: setattr(w.active["A2"], "value", "x" * (MAX_FIELD + 1)),
        lambda w: setattr(w.active.row_dimensions[2], "hidden", True),
        lambda w: w.active.merge_cells("A1:B1"),
        lambda w: setattr(w.active["A2"], "hyperlink", "https://example.invalid"),
    ],
)
def test_malformed_excel_rejected(change, tmp_path):
    workbook = Workbook()
    workbook.active.append(("a", "b"))
    workbook.active.append(("safe", "value"))
    change(workbook)
    path = tmp_path / "bad.xlsx"
    workbook.save(path)
    with pytest.raises(SafetyError):
        read_excel(path)


def test_non_utf8_and_size_bounds(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"\xff")
    with pytest.raises(SafetyError):
        read_excel(path)
    path.write_bytes(b"x" * 21)
    with pytest.raises(SafetyError):
        read_bounded(path, 20)


def test_excel_only_and_literal_text_round_trip(tmp_path):
    path = tmp_path / "literal.xlsx"
    path.write_bytes(excel_bytes(Table(("key", "value"), ({"key": "000123", "value": "=1+2"},))))
    assert read_excel(path).rows == ({"key": "000123", "value": "=1+2"},)
    with pytest.raises(SafetyError, match="Only .xlsx"):
        read_excel(tmp_path / "old.csv")


def test_source_formula_uses_saved_result_but_returned_formula_is_rejected(tmp_path):
    workbook = Workbook()
    workbook.active.append(("key", "amount"))
    workbook.active.append(("SYNTH-001", "=1+2"))
    path = tmp_path / "formula.xlsx"
    workbook.save(path)

    with pytest.raises(SafetyError, match="no saved result"):
        read_excel(path, allow_cached_formulas=True)
    with pytest.raises(SafetyError, match="unsupported cell type"):
        read_excel(path)

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    original = BytesIO(path.read_bytes())
    updated = BytesIO()
    with ZipFile(original) as source, ZipFile(updated, "w") as destination:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                tree = ElementTree.fromstring(data)
                cell = tree.find(f".//{{{namespace}}}c[@r='B2']")
                assert cell is not None
                value = cell.find(f"{{{namespace}}}v")
                assert value is not None
                value.text = "3"
                data = ElementTree.tostring(tree, encoding="utf-8")
            destination.writestr(entry, data)
    path.write_bytes(updated.getvalue())

    table = read_excel(path, allow_cached_formulas=True)
    assert table.rows == ({"key": "SYNTH-001", "amount": "3"},)
    assert table.formula_cells == 1
    with pytest.raises(SafetyError, match="unsupported cell type"):
        read_excel(path)


def test_source_dates_are_text_for_review_but_returned_dates_are_rejected(tmp_path):
    workbook = Workbook()
    workbook.active.append(("key", "date", "datetime", "time"))
    workbook.active.append(
        ("SYNTH-001", date(2030, 4, 5), datetime(2030, 4, 5, 9, 30), time(9, 30))
    )
    path = tmp_path / "dates.xlsx"
    workbook.save(path)

    table = read_excel(path, allow_source_dates=True)
    assert table.rows == (
        {
            "key": "SYNTH-001",
            "date": "2030-04-05T00:00:00",
            "datetime": "2030-04-05T09:30:00",
            "time": "09:30:00",
        },
    )
    assert table.date_cells == 3
    with pytest.raises(SafetyError, match="unsupported cell type"):
        read_excel(path)


def test_selected_worksheet_only(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Instructions"
    workbook.active.append(("not_data",))
    workbook.active.append(("=1+2",))
    data = workbook.create_sheet("Allocations")
    data.append(("key", "value"))
    data.append(("000123", "Synthetic"))
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    path = tmp_path / "multiple.xlsx"
    workbook.save(path)

    assert list_excel_sheets(path) == ("Instructions", "Allocations")
    with pytest.raises(SafetyError, match="Select one or more worksheets"):
        read_excel(path)
    assert read_excel(path, "Allocations").rows == ({"key": "000123", "value": "Synthetic"},)
    with pytest.raises(SafetyError, match="missing, hidden or repeated"):
        read_excel(path, "Hidden")


def test_selected_worksheets_append_rows_and_require_matching_headings(tmp_path, monkeypatch):
    workbook = Workbook()
    first = workbook.active
    first.title = "Earlier"
    first.append(("key", "team"))
    first.append(("SYNTH-001", "Robot A"))
    second = workbook.create_sheet("New")
    second.append(("key", "team"))
    second.append(("SYNTH-002", "Robot B"))
    path = tmp_path / "cohorts.xlsx"
    workbook.save(path)
    assert read_excel(path, ("Earlier", "New")).rows == (
        {"key": "SYNTH-001", "team": "Robot A"},
        {"key": "SYNTH-002", "team": "Robot B"},
    )
    with pytest.raises(SafetyError, match="repeated"):
        read_excel(path, ("Earlier", "Earlier"))
    monkeypatch.setattr("safeset.ingestion.MAX_ROWS", 1)
    with pytest.raises(SafetyError, match="combined row limit"):
        read_excel(path, ("Earlier", "New"))
    monkeypatch.setattr("safeset.ingestion.MAX_ROWS", 50_000)
    second["B1"] = "allocation"
    workbook.save(path)
    with pytest.raises(SafetyError, match="identical headings"):
        read_excel(path, ("Earlier", "New"))


def test_structured_tables_use_only_their_ranges(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Synthetic allocation report title"
    sheet["A2"] = "=1+2"  # Unselected note outside either table.
    sheet["H3"] = "Synthetic instructions beside tables"
    sheet["B3"], sheet["C3"] = "key", "team"
    sheet["B4"], sheet["C4"] = "SYNTH-001", "Robot A"
    sheet.add_table(ExcelTable(displayName="EarlierTable", ref="B3:C4"))
    sheet["E3"], sheet["F3"] = "key", "team"
    sheet["E4"], sheet["F4"] = "SYNTH-002", "Robot B"
    sheet.add_table(ExcelTable(displayName="NewTable", ref="E3:F4"))
    path = tmp_path / "tables.xlsx"
    workbook.save(path)
    assert read_excel(path).rows == (
        {"key": "SYNTH-001", "team": "Robot A"},
        {"key": "SYNTH-002", "team": "Robot B"},
    )
    sheet["F4"] = "=1+2"
    workbook.save(path)
    with pytest.raises(SafetyError, match="unsupported cell type"):
        read_excel(path)


def test_structured_table_rejects_header_metadata_mismatch(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("Student Number", "Team"))
    sheet.append(("SYNTH-001", "Robot A"))
    sheet.add_table(ExcelTable(displayName="Allocations", ref="A1:B2"))
    path = tmp_path / "mismatch.xlsx"
    workbook.save(path)

    changed = load_workbook(path)
    changed.active["A1"] = "Different Heading"
    changed.save(path)
    with pytest.raises(SafetyError, match="table headings differ"):
        read_excel(path)


def test_structured_table_rejects_mismatched_headings_and_hidden_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("key", "team", None, "key", "allocation"))
    sheet.append(("SYNTH-001", "Robot A", None, "SYNTH-002", "Robot B"))
    sheet.add_table(ExcelTable(displayName="EarlierTable", ref="A1:B2"))
    sheet.add_table(ExcelTable(displayName="NewTable", ref="D1:E2"))
    path = tmp_path / "tables.xlsx"
    workbook.save(path)
    with pytest.raises(SafetyError, match="identical headings"):
        read_excel(path)
    sheet["E1"] = "team"
    sheet.row_dimensions[2].hidden = True
    workbook.save(path)
    with pytest.raises(SafetyError, match="hidden rows"):
        read_excel(path)


def test_structured_table_skips_totals_row(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("key", "amount"))
    sheet.append(("SYNTH-001", "1"))
    sheet.append(("SYNTH-002", "2"))
    sheet.append(("Total", "=SUM(B2:B3)"))
    sheet.add_table(
        ExcelTable(
            displayName="Amounts",
            ref="A1:B4",
            totalsRowCount=1,
            totalsRowShown=True,
        )
    )
    path = tmp_path / "totals.xlsx"
    workbook.save(path)
    assert read_excel(path).rows == (
        {"key": "SYNTH-001", "amount": "1"},
        {"key": "SYNTH-002", "amount": "2"},
    )


def test_plain_sheet_filter_and_chart_do_not_block_data(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("key", "amount"))
    sheet.append(("SYNTH-001", 1))
    sheet.append(("SYNTH-002", 2))
    sheet.auto_filter.ref = "A1:B3"
    sheet.cell(row=100_000, column=200).number_format = "@"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "D1")
    path = tmp_path / "filtered.xlsx"
    workbook.save(path)
    assert read_excel(path).rows == (
        {"key": "SYNTH-001", "amount": "1"},
        {"key": "SYNTH-002", "amount": "2"},
    )


@pytest.mark.parametrize(
    "cell,error",
    [("DY2", "column limit"), ("A50002", "row limit")],
)
def test_populated_cells_still_enforce_sheet_limits(tmp_path, cell, error):
    workbook = Workbook()
    workbook.active.append(("key",))
    workbook.active.append(("SYNTH-001",))
    workbook.active[cell] = "Invented extra value"
    path = tmp_path / "wide.xlsx"
    workbook.save(path)
    with pytest.raises(SafetyError, match=error):
        read_excel(path)


@pytest.mark.parametrize(
    "text",
    [
        "version: 1\nversion: 1\n",
        "a: &a [*a]",
        "x: !!python/object:thing {}",
        "version: true\ncolumns: {}\nmin_group_size: 2",
        "- a\n- b",
        "version: 3\ncolumns: {}\nmin_group_size: 2",
    ],
)
def test_hostile_or_invalid_yaml(text, tmp_path):
    path = tmp_path / "bad.yaml"
    path.write_text(text)
    with pytest.raises(SafetyError):
        load_policy(path)


@pytest.mark.parametrize(
    "change",
    [
        lambda p: p.update(unexpected=True),
        lambda p: p.update(min_group_size=1),
        lambda p: p.update(min_group_size=True),
        lambda p: p["columns"]["campus"].update(classification="unknown"),
        lambda p: p["columns"]["campus"].update(classification="free_text"),
        lambda p: p["columns"]["campus"].update(action="redact"),
        lambda p: p["columns"]["campus"].update(allowed_values=["x@example.invalid"]),
        lambda p: p["columns"]["campus"].update(allowed_values=["2020-01-01"]),
        lambda p: p["columns"]["campus"].update(allowed_values=["12345678"]),
        lambda p: p["columns"]["campus"].update(allowed_values=["=SUM(1,2)"]),
        lambda p: p["columns"]["campus"].update(
            allowed_values=["this is deliberately long personal prose"]
        ),
        lambda p: p["columns"]["campus"].update(allowed_values=[True]),
        lambda p: p["columns"]["campus"].update(allowed_values=["Moon", "Moon"]),
        lambda p: p["columns"]["gpa"].update(bounds=[0, 0]),
        lambda p: p["columns"]["gpa"].update(bounds=[0, float("inf")]),
        lambda p: p["columns"]["gpa"].update(bounds=[-1, 7]),
        lambda p: p["columns"]["gpa"].update(bounds=[True, 7]),
        lambda p: p["columns"]["gpa"].update(max_decimal_places=True),
        lambda p: p["columns"]["gpa"].update(max_decimal_places=7),
        lambda p: p["columns"]["gpa"].update(bins=[[0, 7]]),
        lambda p: p["columns"]["campus"].update(bins=[[0, 7]]),
        lambda p: p["columns"]["student_number"].update(action="drop"),
        lambda p: p["columns"]["email"].update(action="pseudonymise"),
        lambda p: p["columns"].update(record_id={"action": "drop", "classification": "unknown"}),
        lambda p: p["columns"]["notes"].update(
            action="keep", classification="analytical_attribute", allowed_values=["ok"]
        ),
        lambda p: p["columns"]["email"].update(
            action="keep", classification="analytical_attribute", allowed_values=["ok"]
        ),
    ],
)
def test_policy_fail_closed(change):
    raw = yaml.safe_load((ROOT / "examples/example-policy.yaml").read_text())
    change(raw)
    with pytest.raises(SafetyError):
        parse_policy(raw)


@pytest.mark.parametrize(
    "value",
    ["", "NaN", "Infinity", "-1", "7.1", "secret-nonnumeric", "4.123", "4e0", "+4.2"],
)
def test_invalid_numeric_values_do_not_leak(source, policy, value):
    rows = deepcopy(source.rows)
    rows[0]["gpa"] = value
    with pytest.raises(SafetyError) as caught:
        sanitise(Table(source.columns, rows), policy)
    if value:
        assert value not in str(caught.value)


def test_numeric_value_is_preserved_without_formatting(source, policy):
    rows = tuple({**row, "gpa": "04.20"} for row in source.rows)
    candidate = sanitise(Table(source.columns, rows), policy)
    assert [row["gpa"] for row in candidate.table.rows] == ["4.2"] * len(rows)


def test_unapproved_category_does_not_leak(source, policy):
    rows = deepcopy(source.rows)
    rows[0]["campus"] = "Secret synthetic campus"
    with pytest.raises(SafetyError) as caught:
        sanitise(Table(source.columns, rows), policy)
    assert "Secret synthetic campus" not in str(caught.value)


@pytest.mark.parametrize(
    "value,label",
    [("0", "[0, 4)"), ("4", "[4, 5)"), ("5", "[5, 6)"), ("6", "[6, 7]"), ("7", "[6, 7]")],
)
def test_v1_bin_boundaries(source, value, label):
    raw = yaml.safe_load((ROOT / "examples/example-policy.yaml").read_text())
    raw["version"] = 1
    raw["columns"]["campus"]["action"] = "keep"
    raw["columns"]["subject"]["action"] = "keep"
    raw["columns"]["gpa"] = {
        "action": "bin",
        "classification": "quasi_identifier",
        "bins": [[0, 4], [4, 5], [5, 6], [6, 7]],
    }
    policy = parse_policy(raw)
    rows = tuple({**r, "gpa": value} for r in source.rows)
    assert sanitise(Table(source.columns, rows), policy).table.rows[0]["gpa"] == label


def test_v1_rejects_new_actions():
    raw = yaml.safe_load((ROOT / "examples/example-policy.yaml").read_text())
    raw["version"] = 1
    with pytest.raises(SafetyError):
        parse_policy(raw)


@pytest.mark.parametrize("mode", ["extra", "missing", "duplicate_key", "empty_key", "empty"])
def test_source_schema_and_keys(source, policy, mode):
    rows = deepcopy(source.rows)
    columns = source.columns
    if mode == "extra":
        columns = (*columns, "new_sensitive_field")
        rows = tuple({**r, "new_sensitive_field": "invented"} for r in rows)
    elif mode == "missing":
        columns = columns[:-1]
    elif mode == "duplicate_key":
        rows[0]["student_number"] = rows[1]["student_number"]
    elif mode == "empty_key":
        rows[0]["student_number"] = ""
    else:
        rows = ()
    with pytest.raises(SafetyError):
        sanitise(Table(columns, rows), policy)


def test_fifo_rejected_without_blocking(tmp_path):
    import os

    path = tmp_path / "pipe"
    os.mkfifo(path)
    with pytest.raises(SafetyError, match="regular"):
        read_bounded(path)


@pytest.mark.parametrize("value", ["=1+2", "+1234", "@command", "bad\tkey"])
def test_source_keys_must_be_restorable(source, policy, value):
    rows = tuple(
        {**r, "student_number": value if i == 0 else r["student_number"]}
        for i, r in enumerate(source.rows)
    )
    with pytest.raises(SafetyError):
        sanitise(Table(source.columns, rows), policy)
