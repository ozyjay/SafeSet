"""Synthetic participant additions/removals, exact identity binding and package preservation."""

import copy
import io
import json
import socket
import zipfile

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.table import TableColumn

from safeset.desktop_bridge import Bridge
from safeset.errors import SafetyError
from safeset.ingestion import Table, excel_workbook_bytes
from safeset.reconciliation import approve_participants, prepare_participants, update_participants

from . import test_workbook_editing as editing_support
from .conftest import PASSPHRASE
from .test_desktop_bridge import call
from .test_workbook_editing import change_allocation, returned_tables

editing = editing_support.editing


def config(**kwargs):
    return dict(
        target_sheet="Allocations",
        reference_sheet="Participants",
        additions_sheet="SafeSet additions",
        include_additions=True,
        include_removals=False,
        column_sources={},
        **kwargs,
    )


def prepare(files, options=None):
    return prepare_participants(
        files.returned, files.source, files.bundle, files.restored, PASSPHRASE, options or config()
    )


def add_proposal(files, mutate=None):
    tables = returned_tables(files)
    participant = tables["Participants"].rows[2]
    rows = [
        {
            "entity_id": participant["entity_id"],
            "team": participant["team"],
            "state": "Current",
            "score": "3",
        }
    ]
    if mutate:
        mutate(rows, tables)
    tables["SafeSet additions"] = Table(tuple(rows[0]), tuple(rows))
    files.returned.write_bytes(excel_workbook_bytes(tables))


def approve(review, **kwargs):
    return approve_participants(
        review,
        {s: tuple(c) for s, c in review.summary["changes"].items()},
        bool(review.summary["additions"]),
        bool(review.summary["removals"]),
        authorised=True,
        **kwargs,
    )


def test_missing_assignments_review_without_export(editing):
    review = prepare(editing)
    assert review.summary["available_additions"] == 1
    assert review.summary["missing_assignments"] == 1
    assert not review.summary["ready"]
    assert review.workbook_bytes is None
    with pytest.raises(SafetyError, match="Review every participant"):
        approve(review)
    assert not editing.restored.exists()
    # The copyable prompt has schema, never authenticated identity values.
    prompt = review.summary["analysis_prompt"]
    assert "SafeSet additions" in prompt and "reference" in prompt
    assert all(
        entity not in prompt and identity not in prompt
        for entity, identity in review.bundle["entities"].items()
    )


def test_addition_and_edit_preserve_other_zip_members(editing, monkeypatch):
    monkeypatch.setattr(socket, "socket", lambda *a, **k: pytest.fail("Network attempted"))
    change_allocation(editing)
    add_proposal(editing)
    review = prepare(editing)
    assert review.summary["ready"]
    assert review.summary["target_rows"] == 3
    assert review.summary["changes"]["Allocations"] == {"team": 1, "state": 1, "score": 1}
    assert approve(review) == 3
    book = load_workbook(editing.restored)
    sheet = book["Allocations"]
    assert sheet["B6"].value == "SYNTH-003"
    assert sheet["C6"].value == "Synthetic Beta"
    assert sheet["E6"].value == 3 and sheet["E6"].data_type == "n"
    assert sheet["C4"].value == "Synthetic Beta"
    assert sheet["C4"].fill.fgColor.rgb == "00FFFF00"
    assert sheet.tables["SyntheticAllocations"].ref == "B3:E6"
    assert book["Summary"]["A2"].value == "=SUM(Allocations!E4:E5)"
    book.close()
    with zipfile.ZipFile(editing.source) as original, zipfile.ZipFile(editing.restored) as result:
        changed = {n for n in original.namelist() if original.read(n) != result.read(n)}
        assert changed == {"xl/worksheets/sheet1.xml", "xl/tables/table1.xml", "xl/workbook.xml"}


@pytest.mark.parametrize(
    "mutation",
    [
        lambda rows, tables: rows.append(dict(rows[0])),
        lambda rows, tables: rows[0].update(entity_id="not-an-entity"),
        lambda rows, tables: rows[0].update(entity_id=tables["Allocations"].rows[0]["entity_id"]),
        lambda rows, tables: rows[0].update(record_id="invented"),
        lambda rows, tables: rows[0].update(team="invented-code"),
        lambda rows, tables: rows[0].update(state="invented-category"),
        lambda rows, tables: rows[0].update(score="11"),
        lambda rows, tables: rows[0].update(score="=1+1"),
    ],
)
def test_bad_proposals_fail_closed(editing, mutation):
    add_proposal(editing, mutation)
    with pytest.raises(SafetyError):
        prepare(editing)
    assert not editing.restored.exists()


def test_original_missing_ids_still_rejected(editing):
    add_proposal(editing)
    book = load_workbook(editing.returned)
    book["Allocations"].delete_rows(2)
    book.save(editing.returned)
    with pytest.raises(SafetyError, match="record coverage"):
        prepare(editing)


@pytest.mark.parametrize("which", ["source", "returned"])
def test_stale_review_cannot_publish(editing, which):
    add_proposal(editing)
    review = prepare(editing)
    path = getattr(editing, which)
    book = load_workbook(path)
    book.properties.title = "Synthetic changed metadata"
    book.save(path)
    with pytest.raises(SafetyError, match="changed after"):
        approve(review)
    assert not editing.restored.exists()


def test_explicit_addition_approval_required(editing):
    add_proposal(editing)
    review = prepare(editing)
    with pytest.raises(SafetyError, match="Review every participant"):
        approve_participants(
            review, {"Allocations": (), "Participants": ()}, False, False, authorised=True
        )
    assert not editing.restored.exists()


def test_existing_only_is_explicit_and_can_update_without_secret(editing):
    review = prepare(editing)
    options = config()
    options["include_additions"] = False
    updated = update_participants(review, options)
    assert updated.summary["ready"] and updated.summary["additions"] == 0
    approve(updated)
    assert editing.source.read_bytes() == editing.restored.read_bytes()


def test_bridge_review_tokens_and_value_free_summary(editing):
    bridge = Bridge()
    payload = dict(
        returned=str(editing.returned),
        source=str(editing.source),
        bundle=str(editing.bundle),
        output=str(editing.restored),
        passphrase=PASSPHRASE,
        config=config(),
    )
    first = call(bridge, "prepare_participants", payload)
    assert first["ok"]
    result = first["result"]
    assert not result["ready"]
    assert "SYNTH-" not in json.dumps(result)
    options = config()
    options["include_additions"] = False
    updated = call(
        bridge, "update_participants", {"review_id": result["review_id"], "config": options}
    )
    assert updated["ok"] and updated["result"]["ready"]
    assert result["review_id"] != updated["result"]["review_id"]
    done = call(
        bridge,
        "approve_participants",
        {
            "review_id": updated["result"]["review_id"],
            "approved_changes": {"Allocations": [], "Participants": []},
            "approve_additions": False,
            "approve_removals": False,
        },
    )
    assert done["ok"]
    again = call(
        bridge,
        "approve_participants",
        {
            "review_id": updated["result"]["review_id"],
            "approved_changes": {},
            "approve_additions": False,
            "approve_removals": False,
        },
    )
    assert not again["ok"]


def test_removal_and_local_mapping_with_authenticated_synthetic_bundle(editing):
    # Produce a separate real protection run with a departing synthetic entity and a local field.
    from safeset.desktop_flow import approve_relational_protection, prepare_relational_protection
    from safeset.policy_authoring import RuleDraft

    book = load_workbook(editing.source)
    book["Participants"].delete_rows(3)  # SYNTH-002 absent; SYNTH-003 still the newcomer.
    book["Allocations"]["F3"] = "local_note"
    book["Allocations"]["F4"] = "Synthetic original note A"
    book["Allocations"]["F5"] = "Synthetic original note B"
    book["Allocations"].tables["SyntheticAllocations"].ref = "B3:F5"
    book["Allocations"].tables["SyntheticAllocations"].tableColumns.append(
        TableColumn(id=5, name="local_note")
    )
    book["Participants"]["C1"] = "local_note"
    book["Participants"]["C2"] = "Synthetic reference note A"
    book["Participants"]["C3"] = "Synthetic newcomer note"
    book.save(editing.source)
    drafts = copy.deepcopy(editing.drafts)
    for sheet in drafts:
        drafts[sheet]["local_note"] = RuleDraft("drop", "direct_identifier")
    new_bundle = editing.bundle.with_name("synthetic-new.enc")
    new_protected = editing.protected.with_name("synthetic-new.xlsx")
    protection = prepare_relational_protection(
        editing.source,
        tuple(drafts),
        drafts,
        "2",
        new_protected,
        new_bundle,
        "controlled_pseudonymisation",
        ("team",),
        editing.fields,
    )
    approve_relational_protection(protection, PASSPHRASE, approved=True)
    editing.bundle = new_bundle
    editing.returned.write_bytes(new_protected.read_bytes())
    tables = returned_tables(editing)
    person = tables["Participants"].rows[1]
    tables["SafeSet additions"] = Table(
        ("entity_id", "team", "state", "score"),
        (
            {
                "entity_id": person["entity_id"],
                "team": person["team"],
                "state": "Current",
                "score": "5",
            },
        ),
    )
    editing.returned.write_bytes(excel_workbook_bytes(tables))
    options = config()
    options["include_removals"] = True
    review = prepare(editing, options)
    assert review.summary["missing_fields"] == ["local_note"]
    assert not review.summary["ready"]
    options["column_sources"] = {"local_note": "local_note"}
    ready = update_participants(review, options)
    assert ready.summary["additions"] == ready.summary["removals"] == 1
    with pytest.raises(SafetyError):
        approve_participants(
            ready, {"Allocations": (), "Participants": ()}, True, False, authorised=True
        )
    approve(ready)
    result = load_workbook(editing.restored)
    assert result["Allocations"]["B5"].value == "SYNTH-003"
    assert result["Allocations"]["F5"].value == "Synthetic newcomer note"
    assert result["Allocations"]["F4"].value == "Synthetic original note A"
    assert result["Participants"]["C3"].value == "Synthetic newcomer note"
    result.close()


def test_writer_rejects_collision_below_target(editing):
    from safeset.ingestion import read_excel_sheets
    from safeset.participant_workbook import replace_participants

    source = read_excel_sheets(editing.source, ("Allocations",))["Allocations"]
    book = load_workbook(editing.source)
    book["Allocations"]["B6"] = "Synthetic footer must survive"
    stream = io.BytesIO()
    book.save(stream)
    proposed = Table(source.columns, (*source.rows, source.rows[0]))
    with pytest.raises(SafetyError, match="safe space"):
        replace_participants(
            stream.getvalue(), editing.source, "Allocations", proposed, [0, 1, None], {"score"}
        )


@pytest.mark.parametrize("kind", ["extra_sheet", "hidden_sheet", "reference_edit"])
def test_reconciliation_preserves_reference_and_sheet_boundaries(editing, kind):
    add_proposal(editing)
    book = load_workbook(editing.returned)
    if kind == "reference_edit":
        book["Participants"]["C2"] = "invented-code"
    else:
        extra = book.create_sheet("Synthetic unexpected")
        extra.append(("field",))
        extra.append(("value",))
        if kind == "hidden_sheet":
            extra.sheet_state = "hidden"
    book.save(editing.returned)
    with pytest.raises(SafetyError):
        prepare(editing)
    assert not editing.restored.exists()


@pytest.mark.parametrize(
    "field,value",
    [
        ("reference_sheet", "Allocations"),
        ("target_sheet", "Participants"),
        ("include_removals", "true"),
        ("column_sources", {"student_key": "student_key"}),
        ("column_sources", {"team": "team"}),
    ],
)
def test_reconciliation_configuration_cannot_override_permissions(editing, field, value):
    options = config()
    options[field] = value
    with pytest.raises(SafetyError):
        prepare(editing, options)


def test_removal_compacts_surviving_cells_without_deleting_worksheet_rows(editing):
    from safeset.ingestion import read_excel_sheets
    from safeset.participant_workbook import replace_participants

    original = read_excel_sheets(editing.source, ("Allocations",))["Allocations"]
    book = load_workbook(editing.source)
    book["Allocations"]["G5"] = "Synthetic neighbouring cell"
    data = io.BytesIO()
    book.save(data)
    result = replace_participants(
        data.getvalue(),
        editing.source,
        "Allocations",
        Table(original.columns, (original.rows[1],)),
        [1],
        {"score"},
    )
    restored = load_workbook(io.BytesIO(result))
    assert restored["Allocations"]["B4"].value == "SYNTH-002"
    assert restored["Allocations"]["E4"].value == 2
    assert restored["Allocations"]["B5"].value is None
    assert restored["Allocations"]["G5"].value == "Synthetic neighbouring cell"
    assert restored["Allocations"].tables["SyntheticAllocations"].ref == "B3:E4"
    restored.close()
