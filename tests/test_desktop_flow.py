"""Safety checks for the desktop controller without requiring a display server."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table as ExcelTable

from safeset.desktop import ACTION_LABELS, Desktop
from safeset.desktop_flow import approve_export, inspect_returned, prepare_export, restore_results
from safeset.errors import SafetyError
from safeset.ingestion import Table, excel_bytes, read_excel

from .conftest import PASSPHRASE, ROOT


@pytest.mark.parametrize(
    "action,enabled",
    [
        (None, False),
        ("drop", False),
        ("pseudonymise", False),
        ("keep", True),
        ("code", True),
        ("bin", True),
        ("keep_numeric", True),
    ],
)
def test_policy_settings_button_follows_action(action, enabled):
    class Selection:
        def get(self):
            return ACTION_LABELS[action] if action else ""

    class Button:
        def configure(self, **kwargs):
            self.text = kwargs["text"]

        def state(self, flags):
            self.flags = flags

    desktop = Desktop.__new__(Desktop)
    button = Button()
    desktop.policy_controls = {"Invented Field": (Selection(), None)}
    desktop.policy_details_buttons = {"Invented Field": button}
    desktop._sync_policy_settings_button("Invented Field")
    assert button.text == "Settings…"
    assert button.flags == (["!disabled"] if enabled else ["disabled"])


def test_review_does_not_publish_and_requires_approval(destinations):
    output, mapping = destinations
    review = prepare_export(
        ROOT / "examples/synthetic_students.xlsx",
        ROOT / "examples/example-policy.yaml",
        output,
        mapping,
    )
    assert review.validation.passed
    assert not output.exists() and not mapping.exists()
    with pytest.raises(SafetyError, match="approval"):
        approve_export(review, PASSPHRASE, approved=False)
    assert not output.exists() and not mapping.exists()


def test_prepare_export_from_selected_sheet(destinations, tmp_path):
    source = tmp_path / "multiple.xlsx"
    workbook = load_workbook(ROOT / "examples/synthetic_students.xlsx")
    workbook.active.title = "Allocations"
    workbook.create_sheet("Notes").append(("Invented instructions",))
    workbook.save(source)
    output, mapping = destinations
    with pytest.raises(SafetyError, match="Select one or more worksheets"):
        prepare_export(source, ROOT / "examples/example-policy.yaml", output, mapping)
    review = prepare_export(
        source, ROOT / "examples/example-policy.yaml", output, mapping, "Allocations"
    )
    assert review.validation.passed
    assert review.sheet == "Allocations"
    assert not output.exists() and not mapping.exists()


def test_overlapping_participants_across_sheets_block_export(destinations, tmp_path):
    source = tmp_path / "overlap.xlsx"
    workbook = load_workbook(ROOT / "examples/synthetic_students.xlsx")
    workbook.active.title = "Earlier"
    workbook.copy_worksheet(workbook.active).title = "Updated"
    workbook.save(source)
    output, mapping = destinations
    with pytest.raises(SafetyError, match="non-empty and unique"):
        prepare_export(
            source,
            ROOT / "examples/example-policy.yaml",
            output,
            mapping,
            ("Earlier", "Updated"),
        )
    assert not output.exists() and not mapping.exists()


def test_combined_worksheets_export_once(destinations, tmp_path):
    source = tmp_path / "cohorts.xlsx"
    workbook = load_workbook(ROOT / "examples/synthetic_students.xlsx")
    earlier = workbook.active
    earlier.title = "Earlier"
    updated = workbook.copy_worksheet(earlier)
    updated.title = "Updated"
    earlier.delete_rows(4, 2)
    updated.delete_rows(2, 2)
    workbook.save(source)
    output, mapping = destinations
    review = prepare_export(
        source,
        ROOT / "examples/example-policy.yaml",
        output,
        mapping,
        ("Earlier", "Updated"),
    )
    assert review.validation.passed
    assert review.source_rows == 4
    approve_export(review, PASSPHRASE, approved=True)
    assert len(read_excel(output).rows) == 4


def test_structured_table_can_be_sanitised(destinations, tmp_path):
    source = tmp_path / "structured.xlsx"
    workbook = load_workbook(ROOT / "examples/synthetic_students.xlsx")
    workbook.active.move_range("A1:G5", rows=2, cols=1)
    workbook.active["A1"] = "Invented private note outside the table"
    workbook.active.add_table(ExcelTable(displayName="Participants", ref="B3:H7"))
    workbook.save(source)
    output, mapping = destinations
    review = prepare_export(source, ROOT / "examples/example-policy.yaml", output, mapping)
    assert review.validation.passed
    approve_export(review, PASSPHRASE, approved=True)
    assert len(read_excel(output).rows) == 4
    assert "Invented private note" not in str(read_excel(output).rows)


def test_failed_review_cannot_export(destinations, tmp_path):
    output, mapping = destinations
    source = tmp_path / "synthetic-invalid.xlsx"
    original = read_excel(ROOT / "examples/synthetic_students.xlsx")
    rows = tuple(
        {**row, "gpa": "5.2"} if index == 0 else row for index, row in enumerate(original.rows)
    )
    source.write_bytes(excel_bytes(Table(original.columns, rows)))
    review = prepare_export(source, ROOT / "examples/example-policy.yaml", output, mapping)
    assert not review.validation.passed
    with pytest.raises(SafetyError, match="Validation failed"):
        approve_export(review, PASSPHRASE, approved=True)
    assert not output.exists() and not mapping.exists()


def test_desktop_flow_synthetic_round_trip(destinations):
    output, mapping = destinations
    review = prepare_export(
        ROOT / "examples/synthetic_students.xlsx",
        ROOT / "examples/example-policy.yaml",
        output,
        mapping,
    )
    approve_export(review, PASSPHRASE, approved=True)
    assert output.exists() and mapping.exists()
    restored = output.parent.parent / "private/restored.xlsx"
    columns = ("campus", "subject", "gpa")
    with pytest.raises(SafetyError, match="authorisation"):
        restore_results(output, mapping, restored, columns, PASSPHRASE, authorised=False)
    assert not restored.exists()
    assert restore_results(output, mapping, restored, columns, PASSPHRASE, authorised=True) == 4
    assert read_excel(restored).rows[0]["student_number"] == "SYNTH-001"


def test_restore_from_selected_result_sheet(destinations):
    output, mapping = destinations
    review = prepare_export(
        ROOT / "examples/synthetic_students.xlsx",
        ROOT / "examples/example-policy.yaml",
        output,
        mapping,
    )
    approve_export(review, PASSPHRASE, approved=True)
    workbook = load_workbook(output)
    workbook.active.title = "Results"
    workbook.create_sheet("Instructions").append(("Invented help",))
    analysed = output.parent / "analysed.xlsx"
    workbook.save(analysed)
    returned = inspect_returned(analysed, "Results")
    assert returned.rows == 4
    restored = output.parent.parent / "private/restored.xlsx"
    assert (
        restore_results(
            analysed,
            mapping,
            restored,
            returned.result_columns,
            PASSPHRASE,
            authorised=True,
            sheet="Results",
        )
        == 4
    )
    assert read_excel(restored).rows[0]["student_number"] == "SYNTH-001"


def test_map_must_be_separate_before_review(destinations):
    output, _ = destinations
    with pytest.raises(SafetyError, match="separate"):
        prepare_export(
            ROOT / "examples/synthetic_students.xlsx",
            ROOT / "examples/example-policy.yaml",
            output,
            Path(output.parent / "map.enc"),
        )
    assert not output.exists()


def test_returned_review_shows_headings_and_counts_only(candidate, tmp_path):
    returned = Table(
        ("record_id", "team"),
        tuple(
            {"record_id": row["record_id"], "team": "Invented Team"} for row in candidate.table.rows
        ),
    )
    path = tmp_path / "returned.xlsx"
    path.write_bytes(excel_bytes(returned))
    review = inspect_returned(path)
    assert review.rows == 4
    assert review.result_columns == ("team",)
    assert "Invented Team" not in str(review)


def test_returned_review_accepts_literal_result_heading(candidate, tmp_path):
    returned = Table(
        ("record_id", "Allocation Result"),
        tuple(
            {"record_id": row["record_id"], "Allocation Result": "Invented Team"}
            for row in candidate.table.rows
        ),
    )
    path = tmp_path / "returned.xlsx"
    path.write_bytes(excel_bytes(returned))
    assert inspect_returned(path).result_columns == ("Allocation Result",)


def test_returned_review_rejects_bad_ids_before_passphrase(candidate, tmp_path):
    returned = Table(
        ("record_id", "team"),
        tuple(
            {"record_id": "private-invalid-id" if i == 0 else row["record_id"], "team": "A"}
            for i, row in enumerate(candidate.table.rows)
        ),
    )
    path = tmp_path / "returned.xlsx"
    path.write_bytes(excel_bytes(returned))
    with pytest.raises(SafetyError) as caught:
        inspect_returned(path)
    assert "private-invalid-id" not in str(caught.value)
