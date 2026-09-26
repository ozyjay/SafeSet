"""Bounded, explicit data-region selection for research workbooks."""

from __future__ import annotations

import io
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .errors import SafetyError
from .ingestion import (
    MAX_COLUMNS,
    MAX_ROWS,
    Table,
    _check_archive,
    _overlaps,
    _read_region,
    read_bounded,
    valid_heading,
)


def _open(path: Path):
    if path.suffix.lower() != ".xlsx":
        raise SafetyError("Only .xlsx Excel workbooks are supported.")
    data = read_bounded(path)
    _check_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False)
        cached = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise SafetyError("Input is not a supported Excel workbook.") from None
    if workbook._external_links:
        raise SafetyError("Excel workbook contains external links.")
    return workbook, cached


def _ref(bounds: tuple[int, int, int, int]) -> str:
    left, top, right, bottom = bounds
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"


def validate_region_selections(selections: object, names: tuple[str, ...]) -> dict:
    if (
        not isinstance(selections, dict)
        or not 1 <= len(selections) <= 128
        or set(selections) != set(names)
    ):
        raise SafetyError("Workbook region selection is invalid.")
    occupied: dict[str, list[tuple[int, int, int, int]]] = {}
    for logical_name, selection in selections.items():
        if (
            not valid_heading(logical_name)
            or len(logical_name) > 31
            or not isinstance(selection, dict)
            or set(selection) != {"sheet", "range"}
            or not isinstance(selection["sheet"], str)
            or not selection["sheet"]
            or not isinstance(selection["range"], str)
        ):
            raise SafetyError("Workbook region selection is invalid.")
        try:
            bounds = range_boundaries(selection["range"])
        except (TypeError, ValueError):
            raise SafetyError("Workbook region selection is invalid.") from None
        if None in bounds or _ref(bounds) != selection["range"]:
            raise SafetyError("Workbook region selection is invalid.")
        if (
            bounds[0] < 1 or bounds[1] < 1
            or bounds[2] - bounds[0] + 1 > MAX_COLUMNS
            or bounds[3] - bounds[1] > MAX_ROWS
        ):
            raise SafetyError("Workbook region selection is invalid.")
        sheet = selection["sheet"]
        if any(_overlaps(bounds, old) for old in occupied.get(sheet, [])):
            raise SafetyError("Selected workbook regions overlap.")
        occupied.setdefault(sheet, []).append(bounds)
    return selections


def discover_regions(path: Path) -> tuple[dict, ...]:
    """Suggest bounded rectangles; suggestions have no authority until confirmed."""
    workbook, _ = _open(path)
    proposals = []
    try:
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            covered = []
            for structured in worksheet.tables.values():
                bounds = range_boundaries(structured.ref)
                covered.append(bounds)
                proposals.append({
                    "sheet": worksheet.title,
                    "range": _ref(bounds),
                    "header_row": bounds[1],
                    "kind": "excel_table",
                })
            occupied = [
                cell for cell in worksheet._cells.values()
                if cell.value is not None
                and not any(
                    left <= cell.column <= right and top <= cell.row <= bottom
                    for left, top, right, bottom in covered
                )
            ]
            if not occupied:
                continue
            min_row = min(cell.row for cell in occupied)
            max_row = max(cell.row for cell in occupied)
            if max_row > MAX_ROWS + 512:
                raise SafetyError("Excel worksheet exceeds the supported row limit.")
            row_columns = defaultdict(set)
            for cell in occupied:
                row_columns[cell.row].add(cell.column)
            groups = []
            start = None
            for row in range(min_row, max_row + 2):
                if row <= max_row and row_columns[row]:
                    if start is None:
                        start = row
                elif start is not None:
                    groups.append((start, row - 1))
                    start = None
            for first, last in groups:
                if first == last:
                    continue
                columns = set().union(*(row_columns[row] for row in range(first, last + 1)))
                if len(columns) < 2 or max(columns) - min(columns) + 1 > MAX_COLUMNS:
                    continue
                bounds = (min(columns), first, max(columns), last)
                if any(_overlaps(bounds, table_bounds) for table_bounds in covered):
                    continue
                proposals.append({
                    "sheet": worksheet.title,
                    "range": _ref(bounds),
                    "header_row": first,
                    "kind": "plain_range",
                })
    finally:
        workbook.close()
    return tuple(proposals)


def read_confirmed_regions(
    path: Path,
    selections: dict[str, dict[str, str]],
    *,
    observe_cell=None,
) -> dict[str, Table]:
    """Read only named, confirmed regions; each name becomes a protected sheet."""
    if not isinstance(selections, dict) or not 1 <= len(selections) <= 128:
        raise SafetyError("Select one or more explicit workbook regions.")
    validate_region_selections(selections, tuple(selections))
    workbook, cached = _open(path)
    results = {}
    occupied: dict[str, list[tuple[int, int, int, int]]] = {}
    try:
        for logical_name, selection in selections.items():
            if (
                selection["sheet"] not in workbook
                or workbook[selection["sheet"]].sheet_state != "visible"
            ):
                raise SafetyError("Workbook region selection is invalid.")
            bounds = range_boundaries(selection["range"])
            sheet = selection["sheet"]
            if any(_overlaps(bounds, old) for old in occupied.get(sheet, [])):
                raise SafetyError("Selected workbook regions overlap.")
            occupied.setdefault(sheet, []).append(bounds)
            matching_table = None
            for structured in workbook[sheet].tables.values():
                table_bounds = range_boundaries(structured.ref)
                if _overlaps(bounds, table_bounds):
                    if table_bounds != bounds or matching_table is not None:
                        raise SafetyError("Select the complete Excel Table region.")
                    matching_table = structured
            last_data_row = bounds[3]
            if matching_table is not None:
                if (
                    matching_table.headerRowCount != 1
                    or matching_table.connectionId is not None
                    or matching_table.tableType not in {None, "worksheet"}
                    or (matching_table.totalsRowCount or 0) not in {0, 1}
                ):
                    raise SafetyError("Excel table has unsupported structure.")
                last_data_row -= matching_table.totalsRowCount or 0
            results[logical_name] = _read_region(
                workbook[sheet], bounds,
                last_data_row=last_data_row,
                cached_worksheet=cached[sheet],
                allow_source_dates=True,
                observe_cell=(
                    (lambda _sheet, coordinate, value, logical=logical_name:
                     observe_cell(logical, coordinate, value))
                    if observe_cell is not None else None
                ),
            )
            if matching_table is not None and matching_table.tableColumns:
                expected = tuple(column.name for column in matching_table.tableColumns)
                if results[logical_name].columns != expected:
                    raise SafetyError("Excel table headings differ from its metadata.")
        if sum(len(table.rows) for table in results.values()) > MAX_ROWS:
            raise SafetyError("Selected workbook regions exceed the combined row limit.")
        return results
    finally:
        workbook.close()
        cached.close()
