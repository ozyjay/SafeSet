"""Bounded Excel ingestion preserving text identifiers."""

import io
import stat
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries

from .errors import SafetyError

MAX_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 128
MAX_FIELD = 4096
MAX_TABLES = 128
CellObserver = Callable[[str, str, str], None]


def valid_heading(name: object) -> bool:
    """Accept a bounded, unambiguous Excel heading without changing its spelling."""
    return (
        isinstance(name, str)
        and bool(name)
        and name == name.strip()
        and len(name) <= 64
        and not any(ord(ch) < 32 or ord(ch) == 127 for ch in name)
    )


@dataclass(frozen=True)
class Table:
    columns: tuple[str, ...]
    rows: tuple[dict[str, str], ...]
    formula_cells: int = 0
    date_cells: int = 0


def _same_headings(
    actual: tuple[str, ...], expected: tuple[str, ...], allow_reordered: bool
) -> bool:
    return set(actual) == set(expected) if allow_reordered else actual == expected


def read_bounded(path: Path, limit: int = MAX_BYTES) -> bytes:
    try:
        if not stat.S_ISREG(path.stat().st_mode):
            raise SafetyError("Input must be a regular local file.")
        with path.open("rb") as stream:
            data = stream.read(limit + 1)
    except OSError:
        raise SafetyError("Unable to read the requested local file.") from None
    if len(data) > limit:
        raise SafetyError("Input exceeds the supported size limit.")
    return data


def _check_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                not entries
                or len(entries) > 1000
                or sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED
                or any(entry.flag_bits & 1 for entry in entries)
            ):
                raise SafetyError("Excel workbook exceeds supported archive limits.")
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)) or any(
                name.startswith("/") or ".." in Path(name).parts for name in names
            ):
                raise SafetyError("Excel workbook has an unsafe archive structure.")
            forbidden = ("vba", "externallinks/", "embeddings/", "connections", "querytables/")
            if any(any(part in name.lower() for part in forbidden) for name in names):
                raise SafetyError("Excel workbook contains unsupported active content.")
    except (zipfile.BadZipFile, OSError, ValueError):
        raise SafetyError("Input is not a supported Excel workbook.") from None


def _cell_text(cell, *, allow_dates: bool = False) -> str:
    value = cell.value
    if value is None:
        return ""
    if allow_dates and cell.data_type in {"d", "n"} and isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if cell.data_type not in {"s", "n", "inlineStr"} or isinstance(
        value, (bool, date, datetime, time)
    ):
        raise SafetyError("Excel workbook contains an unsupported cell type.")
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float)):
        return str(value)
    raise SafetyError("Excel workbook contains an unsupported cell type.")


def list_excel_sheets(path: Path, *, reject_hidden: bool = False) -> tuple[str, ...]:
    """List visible worksheet names without reading cell values into the UI.

    Restoration callers can reject hidden worksheets so an unreviewed sheet cannot
    sit outside the explicit added-analysis worksheet approval flow.
    """
    if path.suffix.lower() != ".xlsx":
        raise SafetyError("Only .xlsx Excel workbooks are supported.")
    data = read_bounded(path)
    _check_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        if workbook._external_links:
            raise SafetyError("Excel workbook contains external links.")
        hidden = tuple(
            sheet.title for sheet in workbook.worksheets if sheet.sheet_state != "visible"
        )
        if reject_hidden and hidden:
            raise SafetyError("Returned workbook contains hidden worksheets.")
        names = tuple(
            sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"
        )
        if not names:
            raise SafetyError("Excel workbook has no visible worksheet.")
        return names
    except SafetyError:
        raise
    except Exception:
        raise SafetyError("Input is not a supported Excel workbook.") from None


def _overlaps(left: tuple[int, int, int, int], right: tuple[int, int, int, int]) -> bool:
    a1, b1, a2, b2 = left
    c1, d1, c2, d2 = right
    return a1 <= c2 and c1 <= a2 and b1 <= d2 and d1 <= b2


def _read_region(
    worksheet,
    bounds: tuple[int, int, int, int],
    *,
    last_data_row: int | None = None,
    cached_worksheet=None,
    allow_source_dates: bool = False,
    allow_cached_formula_blanks: bool = False,
    observe_cell: CellObserver | None = None,
) -> Table:
    min_col, min_row, max_col, max_row = bounds
    if (
        min_col < 1
        or min_row < 1
        or max_col < min_col
        or max_row < min_row
        or max_col - min_col + 1 > MAX_COLUMNS
        or (last_data_row if last_data_row is not None else max_row) - min_row > MAX_ROWS
    ):
        raise SafetyError("Excel table has unsupported dimensions.")
    if any(
        _overlaps(bounds, range_boundaries(str(merged))) for merged in worksheet.merged_cells.ranges
    ):
        raise SafetyError("Excel data range contains merged cells.")
    if any(
        dimension.hidden and min_row <= index <= max_row
        for index, dimension in worksheet.row_dimensions.items()
    ):
        raise SafetyError("Excel data range contains hidden rows.")
    for key, dimension in worksheet.column_dimensions.items():
        index = column_index_from_string(key)
        first = dimension.min or index
        last = dimension.max or index
        if dimension.hidden and first <= max_col and min_col <= last:
            raise SafetyError("Excel data range contains hidden columns.")
    header_cells = next(
        worksheet.iter_rows(min_row=min_row, max_row=min_row, min_col=min_col, max_col=max_col)
    )
    header = tuple(_cell_text(cell) for cell in header_cells)
    if not header or any(not name for name in header):
        raise SafetyError("Excel headings are missing.")
    if len(set(header)) != len(header):
        raise SafetyError("Excel headings are duplicated.")
    if any(not valid_heading(name) for name in header):
        raise SafetyError("Excel headings contain unsupported text.")
    if any(cell.hyperlink or cell.comment for cell in header_cells):
        raise SafetyError("Excel workbook contains unsupported cell features.")
    rows = []
    formula_cells = 0
    date_cells = 0
    data_end = last_data_row if last_data_row is not None else max_row
    for cells in (
        worksheet.iter_rows(
            min_row=min_row + 1,
            max_row=data_end,
            min_col=min_col,
            max_col=max_col,
        )
        if data_end > min_row
        else ()
    ):
        if any(cell.hyperlink or cell.comment for cell in cells):
            raise SafetyError("Excel workbook contains unsupported cell features.")
        values = []
        for cell in cells:
            if cell.data_type == "f" and cached_worksheet is not None:
                value_cell = cached_worksheet[cell.coordinate]
                if value_cell.value is None and not (
                    allow_cached_formula_blanks and value_cell.data_type == "str"
                ):
                    raise SafetyError(
                        "Excel formula has no saved result. Recalculate and save locally."
                    )
                formula_cells += 1
            else:
                value_cell = cell
            value = _cell_text(value_cell, allow_dates=allow_source_dates)
            values.append(value)
            if observe_cell is not None:
                observe_cell(worksheet.title, cell.coordinate, value)
            if allow_source_dates and isinstance(value_cell.value, (date, datetime, time)):
                date_cells += 1
        values = tuple(values)
        if all(value == "" for value in values) or any(
            len(value) > MAX_FIELD or "\x00" in value for value in values
        ):
            raise SafetyError("Excel row shape or field size is invalid.")
        rows.append(dict(zip(header, values, strict=True)))
    return Table(header, tuple(rows), formula_cells, date_cells)


def _read_worksheet(
    worksheet,
    cached_worksheet=None,
    *,
    allow_source_dates: bool = False,
    allow_cached_formula_blanks: bool = False,
    allow_reordered_headings: bool = False,
    observe_cell: CellObserver | None = None,
) -> Table:
    if worksheet.tables:
        if len(worksheet.tables) > MAX_TABLES:
            raise SafetyError("Excel worksheet exceeds the supported table count.")
        ranges = []
        for structured in worksheet.tables.values():
            if (
                structured.headerRowCount != 1
                or structured.connectionId is not None
                or structured.tableType not in {None, "worksheet"}
            ):
                raise SafetyError("Excel table has unsupported structure.")
            bounds = range_boundaries(structured.ref)
            if any(_overlaps(bounds, earlier) for earlier, _ in ranges):
                raise SafetyError("Excel tables overlap.")
            totals = structured.totalsRowCount or 0
            if totals not in {0, 1} or (structured.totalsRowShown and not totals):
                raise SafetyError("Excel table has unsupported totals metadata.")
            ranges.append((bounds, structured))
        ranges.sort(key=lambda item: (item[0][1], item[0][0]))
        tables = []
        for bounds, structured in ranges:
            data = _read_region(
                worksheet,
                bounds,
                last_data_row=bounds[3] - (structured.totalsRowCount or 0),
                cached_worksheet=cached_worksheet,
                allow_source_dates=allow_source_dates,
                allow_cached_formula_blanks=allow_cached_formula_blanks,
                observe_cell=observe_cell,
            )
            if (
                structured.tableColumns
                and tuple(column.name for column in structured.tableColumns) != data.columns
            ):
                raise SafetyError("Excel table headings differ from its metadata.")
            tables.append(data)
        header = tables[0].columns
        if any(
            not _same_headings(table.columns, header, allow_reordered_headings)
            for table in tables[1:]
        ):
            if allow_reordered_headings:
                raise SafetyError("Excel tables must have identical headings.")
            raise SafetyError("Excel tables must have identical headings in the same order.")
        if sum(len(table.rows) for table in tables) > MAX_ROWS:
            raise SafetyError("Excel tables exceed the combined row limit.")
        return Table(
            header,
            tuple(row for table in tables for row in table.rows),
            sum(table.formula_cells for table in tables),
            sum(table.date_cells for table in tables),
        )
    occupied = tuple(
        cell
        for cell in worksheet._cells.values()
        if cell.value is not None or cell.hyperlink or cell.comment
    )
    if not occupied:
        raise SafetyError("Excel headings are missing.")
    last_row = max(cell.row for cell in occupied)
    last_column = max(cell.column for cell in occupied)
    if last_row > MAX_ROWS + 1:
        raise SafetyError("Excel worksheet exceeds the supported row limit.")
    if last_column > MAX_COLUMNS:
        raise SafetyError("Excel worksheet exceeds the supported column limit.")
    return _read_region(
        worksheet,
        (1, 1, last_column, last_row),
        cached_worksheet=cached_worksheet,
        allow_source_dates=allow_source_dates,
        allow_cached_formula_blanks=allow_cached_formula_blanks,
        observe_cell=observe_cell,
    )


def read_excel(
    path: Path,
    sheet: str | tuple[str, ...] | None = None,
    *,
    allow_cached_formulas: bool = False,
    allow_source_dates: bool = False,
    allow_cached_formula_blanks: bool = False,
    allow_reordered_headings: bool = False,
    observe_cell: CellObserver | None = None,
) -> Table:
    if path.suffix.lower() != ".xlsx":
        raise SafetyError("Only .xlsx Excel workbooks are supported.")
    data = read_bounded(path)
    _check_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
        cached_workbook = (
            load_workbook(io.BytesIO(data), read_only=False, data_only=True)
            if allow_cached_formulas
            else None
        )
        if workbook._external_links:
            raise SafetyError("Excel workbook contains external links.")
        visible = tuple(ws.title for ws in workbook.worksheets if ws.sheet_state == "visible")
        if sheet is None:
            if len(visible) != 1:
                raise SafetyError("Select one or more worksheets from the Excel workbook.")
            selected = visible
        elif isinstance(sheet, str):
            selected = (sheet,)
        elif isinstance(sheet, tuple):
            selected = sheet
        else:
            raise SafetyError("Worksheet selection is invalid.")
        if (
            not selected
            or len(set(selected)) != len(selected)
            or any(name not in visible for name in selected)
        ):
            raise SafetyError("Selected worksheet is missing, hidden or repeated.")
        tables = [
            _read_worksheet(
                workbook[name],
                cached_workbook[name] if cached_workbook is not None else None,
                allow_source_dates=allow_source_dates,
                allow_cached_formula_blanks=allow_cached_formula_blanks,
                allow_reordered_headings=allow_reordered_headings,
                observe_cell=observe_cell,
            )
            for name in selected
        ]
        header = tables[0].columns
        if any(
            not _same_headings(table.columns, header, allow_reordered_headings)
            for table in tables[1:]
        ):
            if allow_reordered_headings:
                raise SafetyError("Selected worksheets must have identical headings.")
            raise SafetyError("Selected worksheets must have identical headings in the same order.")
        if sum(len(table.rows) for table in tables) > MAX_ROWS:
            raise SafetyError("Selected worksheets exceed the combined row limit.")
        return Table(
            header,
            tuple(row for table in tables for row in table.rows),
            sum(table.formula_cells for table in tables),
            sum(table.date_cells for table in tables),
        )
    except SafetyError:
        raise
    except Exception:
        raise SafetyError("Input is not a supported Excel workbook.") from None


def read_excel_sheets(
    path: Path,
    sheets: tuple[str, ...],
    *,
    allow_cached_formulas: bool = False,
    allow_source_dates: bool = False,
    allow_cached_formula_blanks: bool = False,
    allow_reordered_headings: bool = False,
    observe_cell: CellObserver | None = None,
) -> dict[str, Table]:
    """Read selected worksheets separately, preserving their distinct schemas."""
    if path.suffix.lower() != ".xlsx":
        raise SafetyError("Only .xlsx Excel workbooks are supported.")
    if not sheets or len(sheets) != len(set(sheets)):
        raise SafetyError("Select one or more distinct worksheets.")
    data = read_bounded(path)
    _check_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
        cached_workbook = (
            load_workbook(io.BytesIO(data), read_only=False, data_only=True)
            if allow_cached_formulas
            else None
        )
        if workbook._external_links:
            raise SafetyError("Excel workbook contains external links.")
        visible = {ws.title for ws in workbook.worksheets if ws.sheet_state == "visible"}
        if any(name not in visible for name in sheets):
            raise SafetyError("Selected worksheet is missing or hidden.")
        result = {
            name: _read_worksheet(
                workbook[name],
                cached_workbook[name] if cached_workbook is not None else None,
                allow_source_dates=allow_source_dates,
                allow_cached_formula_blanks=allow_cached_formula_blanks,
                allow_reordered_headings=allow_reordered_headings,
                observe_cell=observe_cell,
            )
            for name in sheets
        }
        if sum(len(table.rows) for table in result.values()) > MAX_ROWS:
            raise SafetyError("Selected worksheets exceed the combined row limit.")
        return result
    except SafetyError:
        raise
    except Exception:
        raise SafetyError("Input is not a supported Excel workbook.") from None


def require_excel_path(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise SafetyError("Excel output must have an .xlsx filename.")


def excel_bytes(table: Table) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SafeSet"
    sheet.append(table.columns)
    for row in table.rows:
        sheet.append(tuple(row[name] for name in table.columns))
    for cells in sheet:
        for cell in cells:
            cell.data_type = "s"
            cell.number_format = "@"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def excel_workbook_bytes(tables: dict[str, Table]) -> bytes:
    """Create one workbook containing the supplied ordered worksheet tables."""
    if not tables or len(tables) > MAX_TABLES or any(not valid_heading(name) for name in tables):
        raise SafetyError("Protected workbook worksheet structure is invalid.")
    if sum(len(table.rows) for table in tables.values()) > MAX_ROWS:
        raise SafetyError("Protected workbook exceeds the combined row limit.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, table in tables.items():
        if not table.columns or len(table.columns) > MAX_COLUMNS:
            raise SafetyError("Protected worksheet schema is invalid.")
        sheet = workbook.create_sheet(name)
        sheet.append(table.columns)
        for row in table.rows:
            if set(row) != set(table.columns):
                raise SafetyError("Protected worksheet contains a malformed row.")
            sheet.append(tuple(row[column] for column in table.columns))
        for cells in sheet:
            for cell in cells:
                cell.data_type = "s"
                cell.number_format = "@"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
